import io
import pytest
from openpyxl import load_workbook
from backend import ai,acceptance,consistency
from backend.normalize import normalize
from test_core import client,H,new,upload,process

TEXT='''Booking Confirmation
BOOKING NUMBER: 1234567890
PORT OF LOADING:       Harbor A / Terminal              ETA: 20 Sep 2026 07:00(ICT)
                        Terminal continuation
PORT OF DISCHARGE:     New York / Maher Terminal, Berth  ETA: 27 Oct 2026 17:00(EST)
                        64
FINAL DESTINATION:     New York, New York,               ETA: 30 Oct 2026 18:00(EST)
                        United States
ESTIMATED CARGO AVAILABILITY AT DESTINATION HUB: 30 Oct 2026 18:00(EST)
INTENDED VGM CUT-OFF: 18 Sep 2026 10:00(ICT)  INTENDED FCL CY CUT-OFF: 19 Sep 2026 01:00(ICT)
INTENDED ESI CUT-OFF: 18 Sep 2026 09:00(ICT)
BOOKING QTY SIZE/TYPE: 1 X 40' Hi-Cube Refrigerated Container'''

def result(fields):
    return ai.DocumentResult(document_type='Booking',classification_evidence='Booking Confirmation',fields=[
        ai.FieldResult(field_id=key,raw_value=value,evidence_text=quote,source_page=1) for key,value,quote in fields],items=[],containers=[])

def test_split_route_cell_enters_excel_without_employee_review(monkeypatch):
    fields=[('port_loading','Harbor A / Terminal Terminal continuation',TEXT),
            ('port_discharge','New York / Maher Terminal, Berth 64',TEXT),
            ('place_delivery','New York, New York, United States',TEXT),
            ('eta','27 Oct 2026 17:00(EST)',TEXT.splitlines()[4])]
    pages=[{'number':1,'text':TEXT}]
    data=ai.ground_source_format(result(fields),pages)
    records,rejected=ai.validated_subset(data,pages)
    assert not rejected
    assert all(acceptance.readable(r) for r in records)
    assert next(r for r in records if r['key']=='port_discharge')['value']['source_parts'][1]['text']=='64'
    cid=client.post('/api/cases',headers=H,json={'name':'Readable Booking'}).json()['case']['id']
    upload(cid,[('booking.txt',TEXT.encode())])
    monkeypatch.setattr(ai,'key',lambda:'mock')
    monkeypatch.setattr(ai,'analyze',lambda *args:(data,records,{'rejected':[]}))
    monkeypatch.setattr(consistency,'run_ai_review',lambda *args:'')
    b,job=process(cid,'openai');assert not job['errors']
    assert all(r['status']=='auto_accepted' for r in b['records'])
    assert not any(c['key'].startswith('reading:') for c in b['checks'])
    export=client.post('/api/cases/'+cid+'/export',headers=H,json={'mode':'Draft'}).json()
    wb=load_workbook(io.BytesIO(client.get(export['download']).content))
    values=list(wb.worksheets[1].iter_rows(values_only=True))
    assert any(r[0]=='port_discharge' and r[5]=='auto_accepted' for r in values)
    assert any(r[0]=='vgm_cut_off' for r in values)
    wb.close()

@pytest.mark.parametrize('value',['New York / Maher Terminal, Berth 65','New York / Maher Terminal, Berth 27','Harbor A / Terminal Terminal continuation'])
def test_split_cell_does_not_accept_wrong_column_or_invented_value(value):
    data=result([('port_discharge',value,TEXT)])
    with pytest.raises(ValueError):ai.validate_result(data,[{'number':1,'text':TEXT}])

def test_split_cell_never_accepts_a_fabricated_source_quote():
    data=result([('port_discharge','New York / Maher Terminal, Berth 64',TEXT+'\nNOT PRINTED')])
    with pytest.raises(ValueError,match='bằng chứng'):ai.validate_result(data,[{'number':1,'text':TEXT}])

def test_booking_dates_and_cutoffs_are_kept_under_their_own_events():
    fields=[('eta','27 Oct 2026 17:00(EST)',TEXT.splitlines()[4]),
            ('eta','30 Oct 2026 18:00(EST)',TEXT.splitlines()[6]),
            ('eta','30 Oct 2026 18:00(EST)',TEXT.splitlines()[8]),
            ('cut_off','18 Sep 2026 10:00(ICT)',TEXT.splitlines()[9]),
            ('cut_off','19 Sep 2026 01:00(ICT)',TEXT.splitlines()[9]),
            ('packages',"1 X 40' Hi-Cube Refrigerated Container",TEXT.splitlines()[11])]
    pages=[{'number':1,'text':TEXT}]
    records,rejected=ai.validated_subset(ai.ground_source_format(result(fields),pages),pages)
    assert not rejected
    keys={r['key'] for r in records}
    assert {'eta','final_eta','cargo_available_at','vgm_cut_off','cy_cut_off','si_cut_off'}<=keys
    assert not {'packages','cut_off'}&keys
    assert [r['raw'] for r in records if r['key']=='eta']==['27 Oct 2026 17:00(EST)']
    wrong=result([('vgm_cut_off','19 Sep 2026 01:00(ICT)',TEXT)])
    with pytest.raises(ValueError,match='nhãn sự kiện'):ai.validate_result(wrong,pages)

def test_blank_optional_bl_followed_by_date_is_not_a_reading_task():
    text='PACKING LIST\nBILL OF LADING NO.: DATED 21-SEP-2026   BUYER LTD'
    data=ai.DocumentResult(document_type='Packing List',classification_evidence='PACKING LIST',fields=[
        ai.FieldResult(field_id='bl_number',raw_value='DATED 21-SEP-2026',source_page=1,evidence_text=text)],items=[],containers=[])
    assert not ai.ground_source_format(data,[{'number':1,'text':text}]).fields

def test_readable_datetime_retains_time_and_zone_and_invalid_dates_stay_uncertain():
    assert normalize('27 Oct 2026 17:00(EST)','date')['text']=='2026-10-27T17:00 (EST)'
    assert not normalize('27 Oct 2026 17:00(EST)','date')['warning']
    assert normalize('27 Oct 2026 17:00(EST)','date')['text']!=normalize('27 Oct 2026 17:00(ICT)','date')['text']
    assert normalize('31 Feb 2026 17:00(EST)','date')['warning']
    assert normalize('27 Oct 2026 25:00(EST)','date')['warning']

def test_additional_terminal_detail_is_stored_without_claiming_a_source_mismatch():
    from backend.validation import port_detail_compatible
    rows=lambda *values:[{'value':{'text':v}} for v in values]
    assert port_detail_compatible(rows('New York / Maher Terminal, Berth 64','NEW YORK,NY,US'))
    assert port_detail_compatible(rows('Ba Ria-Vung Tau(CM-TV) / TCTT','BA RIA-VUNG TAU(CM-TV)'))
    assert not port_detail_compatible(rows('New York / Terminal A','New York / Terminal B','New York'))
    assert not port_detail_compatible(rows('London / Terminal A','London,UK','London,CANADA'))
    assert not port_detail_compatible(rows('New York / Terminal A','Newark,US'))
