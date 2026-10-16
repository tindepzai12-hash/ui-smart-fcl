import io,json,pytest
from openpyxl import load_workbook
from backend import ai,db,consistency,source_locator
from backend.app import bundle
from backend.normalize import normalize
from test_core import client,H,new,upload,process
from fixtures import INVOICE,PACKING,BL


def test_saved_cases_receive_new_comparison_rules_without_overwriting_values():
    from backend.app import refresh_staff_checks
    cid=new();upload(cid,[('inv.txt',INVOICE.encode()),('pl.txt',PACKING.encode()),('bl.txt',BL.encode())]);before,_=process(cid)
    snapshot={r['id']:(r['value'],r['raw'],r['version']) for r in before['records']}
    db.execute("UPDATE checks SET result='Mismatch' WHERE case_id=? AND key='field:gross_weight'",(cid,))
    db.execute('DELETE FROM schema_versions WHERE version=4')
    refresh_staff_checks()
    after=client.get('/api/cases/'+cid).json()
    assert next(c for c in after['checks'] if c['key']=='field:gross_weight')['result']=='Matched'
    assert {r['id']:(r['value'],r['raw'],r['version']) for r in after['records']}==snapshot
    ids=[c['id'] for c in after['checks']]
    refresh_staff_checks()
    assert [c['id'] for c in client.get('/api/cases/'+cid).json()['checks']]==ids

def test_equal_canonical_sources_never_become_ai_exceptions():
    for key,a,b,kind in [('packages','56 CT','56 CTN','quantity'),('declaration_number','307537505830','307537505830','text')]:
        rows=[{'id':str(i),'key':key,'kind':'field','status':'pending','value':normalize(v,kind)} for i,v in enumerate([a,b])]
        check={'detail':{'sources':[{'id':r['id']} for r in rows]}}
        assert not consistency.semantic_difference(check,rows)
        rows[1]['value']=normalize('57 CT' if key=='packages' else '307537505831',kind)
        assert consistency.semantic_difference(check,rows)

def test_piece_count_groups_without_guessing_weight_or_money():
    assert normalize('1,200 PCS','quantity')['number']=='1200'
    assert not normalize('1,200 PCS','quantity')['warning']
    assert normalize('1,200 KG','weight')['warning']
    assert normalize('1,200 USD','number')['warning']

def test_review_sources_updates_excel_status_but_does_not_match_different_values():
    cid=new();upload(cid,[('inv.txt',INVOICE.encode()),('pl.txt',PACKING.encode()),('bl.txt',BL.replace('1 MT','1100 KG').encode())]);b,_=process(cid)
    check=next(c for c in b['checks'] if c['key']=='field:gross_weight')
    ids={s['id'] for s in check['detail']['sources']};rows=[r for r in b['records'] if r['id'] in ids]
    versions={r['id']:r['version'] for r in rows}
    url='/api/checks/'+check['id']+'/review'
    assert client.post(url,headers=H,json={'versions':{},'note':'Đã đối chiếu bản gốc'}).status_code==409
    response=client.post(url,headers=H,json={'versions':versions,'note':'Đã đối chiếu, yêu cầu khách sửa bản PL vì gross khác.'});assert response.status_code==200,response.text
    b=response.json();check=next(c for c in b['checks'] if c['key']=='field:gross_weight')
    assert check['result']=='Mismatch' and check['resolution']
    assert all(r['status']=='verified' for r in b['records'] if r['id'] in ids)
    assert client.post(url,headers=H,json={'versions':versions,'note':'Một xác nhận cũ'}).status_code==409
    result=client.post('/api/cases/'+cid+'/export',headers=H,json={'mode':'Draft'}).json()
    wb=load_workbook(io.BytesIO(client.get(result['download']).content))
    data=[list(row) for row in wb.worksheets[1].iter_rows(values_only=True)]
    assert any('verified' in row and 'gross_weight' in row for row in data);wb.close()

def rejected_fixture():
    cid=new();upload(cid,[('inv.txt',INVOICE.encode())]);b,_=process(cid);doc=b['documents'][0]
    entry={'field_label':'Tổng trị giá hóa đơn','field_id':'invoice_total','candidate_kind':'field','candidate_raw':{'raw_value':'99999'},
           'attempted_value':'99999','page':1,'evidence':'Total: 1250.25 USD','rejection_reason':'Giá trị không khớp bằng chứng'}
    db.execute('UPDATE documents SET error=?,reading_diagnostics=? WHERE id=?',('AI chưa xác minh được dữ liệu',db.encode([entry]),doc['id']))
    from backend.app import revalidate
    with db.connect() as c:revalidate(c,cid)
    return cid,doc,entry

def test_rejected_attempt_requires_staff_value_token_and_note_then_exports_verified():
    cid,doc,entry=rejected_fixture();url=f"/api/documents/{doc['id']}/diagnostics/0/review"
    payload={'token':source_locator.diagnostic_token(entry),'action':'accept','value':'1250.25 USD','note':'Đã kiểm Total tại trang 1 theo bản gốc'}
    assert client.post(url,headers=H,json={**payload,'token':'stale'}).status_code==409
    assert client.post(url,headers=H,json={**payload,'value':'1,200 USD'}).status_code==422
    response=client.post(url,headers=H,json=payload);assert response.status_code==200,response.text
    b=response.json();manual=[r for r in b['records'] if r['method']=='Manual review']
    assert len(manual)==1 and manual[0]['status']=='verified' and manual[0]['value']['number']=='1250.25'
    assert manual[0]['raw']==entry['candidate_raw']
    assert not any(c['key']=='reading:'+doc['id'] for c in b['checks'])
    assert client.post(url,headers=H,json=payload).status_code==409
    assert not b['documents'][0]['error']

def test_excluding_an_attempt_never_creates_usable_data():
    cid,doc,entry=rejected_fixture();before=len(bundle(cid)['records'])
    r=client.post(f"/api/documents/{doc['id']}/diagnostics/0/review",headers=H,json={'token':source_locator.diagnostic_token(entry),'action':'exclude','note':'Đề xuất sai, không sử dụng giá trị này.'})
    assert r.status_code==200,r.text
    assert len(r.json()['records'])==before
    assert db.rows("SELECT * FROM audit WHERE case_id=?",(cid,))

def test_context_highlight_marks_only_literal_source_tokens():
    words=[{'text':t,'layer':'PDF text','x':i/10,'y':.2,'w':.1,'h':.1,'confidence':None} for i,t in enumerate('ALUMINUM ALLOY CARRIER 600 PCS'.split())]
    assert source_locator.quote_boxes(words,'ALUMINUM ALLOY CARRIER')
    assert not source_locator.quote_boxes(words,'ALUMINUM BRASS CARRIER')

def test_unproven_candidate_still_displays_literal_context_without_server_error(monkeypatch):
    words=[{'text':t,'layer':'PDF text','x':i/10,'y':.2,'w':.1,'h':.1,'confidence':None} for i,t in enumerate('ALUMINUM ALLOY CARRIER'.split())]
    monkeypatch.setattr(source_locator,'layout',lambda *args:words)
    value=source_locator.inspect({'id':'d','path':'unused','name':'source.pdf','ext':'.pdf'},
        {'raw':'invented value','evidence':'ALUMINUM ALLOY CARRIER','method':'OpenAI diagnostic'},
        {'number':1,'method':'PDF text','quality':'good','text':'ALUMINUM ALLOY CARRIER'})
    assert value['boxes'] and all(b['evidence_only'] for b in value['boxes']) and not value['ambiguous']

def test_interrupted_description_keeps_only_literal_commodity_fragment():
    page={'number':1,'text':'INVOICE\nALUMINUM ALLOY CARRIER 600 PCS\nSW-RT73D H325\nHS CODE 871499'}
    result=ai.DocumentResult.model_validate({'document_type':'Invoice','classification_evidence':'INVOICE','fields':[],
       'items':[{'line_number':'1','description':'SW-RT73D H325 ALUMINUM ALLOY CARRIER HS CODE 871499','quantity':'600','unit':'PCS','source_page':1,'evidence_text':'ALUMINUM ALLOY CARRIER 600 PCS'}],'containers':[]})
    grounded=ai.ground_source_format(result,[page]);assert grounded.items[0].description=='ALUMINUM ALLOY CARRIER'
    rows,rejected=ai.validated_subset(grounded,[page]);assert rows and not rejected

def test_warehouse_packages_bind_to_single_printed_row_without_inventing_weight_unit():
    page={'number':1,'text':'YÊU CẦU NHẬP KHO\nUnit Of Measure Quantity Weight Volume\n 1 CT 56 515.080 7.380\n1\nCT\n56\n515.080\n7.380'}
    result=ai.DocumentResult.model_validate({'document_type':'Warehouse Receipt','classification_evidence':'YÊU CẦU NHẬP KHO',
        'fields':[{'field_id':'packages','raw_value':'56','source_page':1,'evidence_text':'Unit Of Measure Quantity Weight Volume'}], 'items':[], 'containers':[]})
    grounded=ai.ground_source_format(result,[page]);field=grounded.fields[0]
    assert field.unit=='CT' and field.evidence_text=='1 CT 56 515.080 7.380'
    rows,rejected=ai.validated_subset(grounded,[page]);assert rows and not rejected
    assert not any(r['key'] in ['gross_weight','cargo_weight','volume'] for r in rows)

def test_item_origin_uses_explicit_origin_header_not_seller_address():
    page={'number':1,'text':'INVOICE\nCountry of origin: VIETNAM\nALUMINUM ALLOY CARRIER 600 PCS'}
    result=ai.DocumentResult.model_validate({'document_type':'Invoice','classification_evidence':'INVOICE','fields':[],
        'items':[{'line_number':'1','description':'ALUMINUM ALLOY CARRIER','quantity':'600','unit':'PCS','origin':'VIETNAM','source_page':1,'evidence_text':'ALUMINUM ALLOY CARRIER 600 PCS'}],'containers':[]})
    grounded=ai.ground_source_format(result,[page]);assert grounded.items[0].cell_evidence[0].evidence_text=='Country of origin: VIETNAM'
    rows,rejected=ai.validated_subset(grounded,[page]);assert rows and not rejected
    grounded=ai.ground_source_format(result,[{**page,'text':page['text'].replace('Country of origin:','Seller address:')}])
    assert not grounded.items[0].cell_evidence

def test_item_code_binds_cells_to_unique_printed_row_despite_wrong_model_pointers():
    page={'number':1,'text':'INVOICE\nITEM-A ALUMINUM ALLOY CARRIER 600 PCS 6.53 3918.00\nITEM-B ALUMINUM ALLOY BRACKET 600 PCS 0.54 324.00',
        'tables':[[['Item No.','Description','Quantity','Unit','Unit Price','Amount'],
                   ['ITEM-A','ALUMINUM ALLOY CARRIER','600','PCS','6.53','3918.00'],
                   ['ITEM-B','ALUMINUM ALLOY BRACKET','600','PCS','0.54','324.00']]]}
    result=ai.DocumentResult.model_validate({'document_type':'Invoice','classification_evidence':'INVOICE','fields':[],
        'items':[{'line_number':'1','product_code':'ITEM-A','description':'ALUMINUM ALLOY CARRIER','quantity':'600','unit':'PCS','unit_price':'6.53','line_amount':'3918.00','source_page':1,'evidence_text':'INVOICE',
            'cell_evidence':[{'attribute':'product_code','evidence_text':'INVOICE'}]}],'containers':[]})
    rows,rejected=ai.validated_subset(ai.ground_source_format(result,[page]),[page]);assert rows and not rejected
    result.items[0].line_amount='324.00'  # Real amount on the other row must not pass.
    with pytest.raises(ValueError,match='line_amount'):
        ai.validated_subset(ai.ground_source_format(result,[page]) ,[page])
    result.items[0].product_code='INVENTED'
    with pytest.raises(ValueError,match='product_code'):
        ai.validated_subset(ai.ground_source_format(result,[page]),[page])
