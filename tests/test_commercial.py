import io,json
from pathlib import Path
import pytest
from openpyxl import load_workbook
from backend import db,extraction,validation,ai
from backend.app import bundle,upgrade_commercial_checklists
from backend.dictionary import default_requirements,PROFILE_TYPES,REVIEW_FIELDS
from test_core import client,H,upload,process

ROOT=Path(__file__).resolve().parents[1]
def create():
 r=client.post('/api/cases',headers=H,json={'name':'DEMO commercial document review','customer':'Demo customer'})
 assert r.status_code==201,r.text
 return r.json()['case']['id']

def sample(name):return (ROOT/'samples'/name).read_bytes()
def prepared(contract='contract.txt',co=False):
 id=create()
 names=['invoice.txt','packing.txt','booking.txt',contract,'arrival.txt']+(['co.txt'] if co else [])
 # Intentionally content-independent names to verify classification uses content.
 upload(id,[(str(i+1)+'.txt',sample(name)) for i,name in enumerate(names)])
 b,j=process(id)
 assert not j['errors'],j
 return id,b

def test_default_profile_and_dictionary():
 id=create();b=bundle(id)
 assert list(b['case']['requirements'])==['Packing List','Booking','Invoice','Contract','Arrival Notice','C/O']
 assert not b['case']['requirements']['C/O']['required']
 d=client.get('/api/dictionary').json()
 assert d['document_types']==PROFILE_TYPES and len(d['review_groups'])==6
 assert not {'hs_code','equipment_count','equipment_type'} & set(d['review_fields'])
 assert not any(c['key']=='document:C/O' and c['result']=='Missing' for c in b['checks'])

@pytest.mark.parametrize('title,kind',[
 ('BOOKING CONFIRMATION','Booking'),('BOOKING NOTE','Booking'),
 ('SALES CONTRACT','Contract'),('SALE CONTRACT','Contract'),
 ('HỢP ĐỒNG MUA BÁN','Contract'),('PURCHASE ORDER','Purchase Order')])
def test_document_headings(title,kind):
 assert extraction.classify([{'text':title+'\nInvoice Reference: ABC'}])[0]==kind

@pytest.mark.parametrize('contract', ['contract.txt','po.txt'])
def test_four_groups_without_co_or_container_can_be_reviewed(contract):
 id,b=prepared(contract)
 assert {d['type'] for d in b['documents']}=={'Invoice','Packing List','Booking','Arrival Notice','Contract' if contract=='contract.txt' else 'Purchase Order'}
 assert not [r for r in b['records'] if r['kind']=='container']
 assert not [c for c in b['checks'] if c['key'].startswith('missing:')],b['checks']
 assert next(c for c in b['checks'] if c['key']=='document:Contract')['result']=='Matched'
 assert next(c for c in b['checks'] if c['key']=='document:C/O')['result']=='Not Applicable'
 for d in b['documents']:
  assert client.patch('/api/documents/'+d['id']+'/classification',headers=H,json={'type':d['type']}).status_code==200
 for r in bundle(id)['records']:
  res=client.post('/api/records/'+r['id']+'/verify',headers=H,json={'status':'verified','version':r['version'],'note':'Đã kiểm bản gốc giả lập'})
  assert res.status_code==200,res.text
 assert not bundle(id)['blockers'],bundle(id)['blockers']
 assert client.post('/api/cases/'+id+'/review',headers=H).status_code==200
 exported=client.post('/api/cases/'+id+'/export',headers=H,json={'mode':'Reviewed'})
 assert exported.status_code==200,exported.text
 wb=load_workbook(io.BytesIO(client.get(exported.json()['download']).content))
 assert len(wb.sheetnames)==6 and 'THONG_TIN_CHUNG_TU' in wb.sheetnames
 assert 'CONTAINER_VAN_CHUYEN' not in wb.sheetnames
 keys={row[0] for row in wb['THONG_TIN_CHUNG_TU'].iter_rows(min_row=2,values_only=True)}
 assert keys<=set(REVIEW_FIELDS) and 'bl_number' in keys
 assert not any(row[1]=='Chuẩn bị khai báo' for row in wb['THIEU_VA_CAN_XAC_MINH'].iter_rows(min_row=2,values_only=True))
 wb.close()

def test_co_present_is_checked_and_po_contract_stay_separate():
 id,b=prepared(co=True)
 upload(id,[('6.txt',sample('po.txt'))]);b,j=process(id)
 assert not j['errors']
 refs={r['key']:r['value']['text'] for r in b['records'] if r['key'] in ['contract_number','po_number','co_number']}
 assert refs['contract_number']=='SC-001' and refs['po_number']=='PO-001' and refs['co_number']=='CO-DEMO-001'
 co=next(d for d in b['documents'] if d['type']=='C/O')
 records=[r for r in b['records'] if not (r['document_id']==co['id'] and r['key']=='origin')]
 checks=validation.build_checks(b['case'],b['documents'],records)
 assert next(c for c in checks if c['key']=='missing:origin')['result']=='Missing'

@pytest.mark.parametrize('co', [False, True])
def test_excel_does_not_call_unextracted_applicable_fields_inapplicable(co):
 id,b=prepared(co=co)
 # Before staff confirms document types, absence is not evidence of inapplicability.
 draft=client.post('/api/cases/'+id+'/export',headers=H,json={'mode':'Draft'})
 assert draft.status_code==200,draft.text
 wb=load_workbook(io.BytesIO(client.get(draft.json()['download']).content))
 rows={row[0]:row for row in wb['THONG_TIN_CHUNG_TU'].iter_rows(min_row=2,values_only=True)}
 assert rows['seller_address'][5]=='Not Yet Checked'
 assert rows['co_form'][5]=='Not Yet Checked'
 wb.close()
 for d in b['documents']:
  assert client.patch('/api/documents/'+d['id']+'/classification',headers=H,json={'type':d['type']}).status_code==200
 date=next(r for r in b['records'] if r['key']=='invoice_date')
 res=client.post('/api/records/'+date['id']+'/verify',headers=H,json={'status':'rejected','version':date['version'],'note':'Test: no accepted invoice date remains'})
 assert res.status_code==200,res.text
 draft=client.post('/api/cases/'+id+'/export',headers=H,json={'mode':'Draft'})
 assert draft.status_code==200,draft.text
 wb=load_workbook(io.BytesIO(client.get(draft.json()['download']).content))
 ws=wb['THONG_TIN_CHUNG_TU']
 rows={row[0]:row for row in ws.iter_rows(min_row=2,values_only=True)}
 assert rows['invoice_date'][5]=='Missing'
 assert rows['seller_address'][5]=='Not Yet Checked'
 assert rows['co_form'][5]==('Not Yet Checked' if co else 'Not Applicable')
 statuses=next(iter(ws.data_validations.dataValidation)).formula1
 assert all(status in statuses for status in ['Not Yet Checked','Not Applicable','Needs Verification'])
 wb.close()

def test_planned_booking_and_multi_delivery_contract_require_scope_review():
 id,b=prepared()
 booking=next(d for d in b['documents'] if d['type']=='Booking')
 records=[{**r,'value':{**r['value'],'number':'1200'}} if r['key']=='gross_weight' and r['document_id']==booking['id'] else r for r in b['records']]
 checks=validation.build_checks(b['case'],b['documents'],records)
 c=next(c for c in checks if c['key']=='field:gross_weight')
 assert c['result']=='Needs Verification' and 'kế hoạch' in c['detail']['reason']
 contract=next(d for d in b['documents'] if d['type']=='Contract')
 records=[{**r,'value':{**r['value'],'number':'150'}} if r['key']=='goods_quantity' and r['document_id']==contract['id'] else r for r in b['records']]
 checks=validation.build_checks(b['case'],b['documents'],records)
 assert next(c for c in checks if c['key']=='field:goods_quantity')['result']=='Needs Verification'

def test_out_of_scope_manual_records_are_rejected():
 id=create()
 for kind,key,value in [('container','container',{'container_number':'MSCU1234567'}),('field','hs_code','12345678')]:
  r=client.post('/api/cases/'+id+'/records',headers=H,json={'kind':kind,'key':key,'value':value,'note':'Demo'})
  assert r.status_code==422,r.text
 assert client.post('/api/cases/'+id+'/records',headers=H,json={'kind':'field','key':'invoice_number','value':'DEMO-001','note':'No source'}).status_code==422

def test_second_invoice_missing_date_is_not_hidden_by_first_invoice():
 id,b=prepared()
 text=sample('invoice.txt').decode('utf-8').replace('000123-A','000124-A')
 upload(id,[('extra.txt','\n'.join(line for line in text.splitlines() if not line.startswith('Invoice Date:')).encode('utf-8'))])
 b,_=process(id)
 c=next(c for c in b['checks'] if c['key']=='missing:invoice_date')
 assert c['result']=='Missing' and 'extra.txt' in c['detail']['reason']

def test_migration_preserves_sources_and_invalidates_old_review():
 id,b=prepared()
 old={'Invoice':{'required':True,'note':'Customer supplied'},'Packing List':{'required':True,'note':'Customer supplied'},'B/L':{'required':True,'note':'Legacy'},'C/O':{'required':True,'note':'Old scope'}}
 db.execute('UPDATE cases SET requirements=?,reviewed_at=? WHERE id=?',(db.encode(old),db.now(),id))
 db.execute('DELETE FROM schema_versions WHERE version=2')
 upgrade_commercial_checklists()
 after=bundle(id)
 assert len(after['documents'])==len(b['documents'])
 assert list(after['case']['requirements'])==list(default_requirements())
 assert not after['case']['requirements']['C/O']['required'] and not after['case']['reviewed_at']
 revision=after['case']['revision'];upgrade_commercial_checklists()
 assert bundle(id)['case']['revision']==revision

def test_ai_booking_evidence_and_po_semantics():
 pages=[{'number':1,'text':'BOOKING CONFIRMATION\nBooking No: BK-001','tables':[]}]
 result=ai.DocumentResult(document_type='Booking',classification_evidence='BOOKING CONFIRMATION',fields=[{'field_id':'booking','raw_value':'BK-001','source_page':1,'evidence_text':'Booking No: BK-001'}],items=[],containers=[])
 assert ai.validate_result(result,pages)[0]['key']=='booking'
 result.fields[0].raw_value='FAKE-NUMBER'
 with pytest.raises(ValueError):ai.validate_result(result,pages)


def test_packing_list_without_own_number_is_not_missing():
 cid=create()
 payload=sample('packing.txt').decode('utf-8')
 payload='\n'.join(line for line in payload.splitlines() if not line.lower().startswith('packing list no:'))
 upload(cid,[('1.txt',payload.encode())])
 b,j=process(cid)
 assert not j['errors']
 assert not any(c['key']=='missing:packing_number' for c in b['checks'])
 assert not any(r['key']=='packing_number' for r in b['records'])

def test_arrival_notice_is_main_but_bill_does_not_supply_booking():
 cid=create()
 upload(cid,[('1.txt',sample('arrival.txt')),('2.txt',sample('bl.txt'))])
 b,j=process(cid)
 assert not j['errors']
 assert [d['type'] for d in b['documents']]==['Arrival Notice']
 assert [d['type'] for d in b['other_files']]==['B/L']
 assert any(r['key']=='bl_number' for r in b['records'])
 assert not any(r['kind']!='field' or r['key']=='booking' for r in b['records'])
 assert next(c for c in b['checks'] if c['key']=='document:Booking')['result']=='Missing'
 assert next(c for c in b['checks'] if c['key']=='document:Arrival Notice')['result']=='Matched'

def test_arrival_scope_migration_retains_file_and_waivers():
 from backend.app import upgrade_arrival_notice_scope
 cid=create()
 upload(cid,[('1.txt',sample('arrival.txt'))]);b,j=process(cid)
 doc=b['documents'][0]
 rules=b['case']['requirements'];rules.pop('Arrival Notice')
 rules['Booking']={'required':False,'note':'Customer waived planned booking'}
 db.execute('UPDATE cases SET requirements=?,reviewed_at=? WHERE id=?',(db.encode(rules),db.now(),cid))
 db.execute("UPDATE documents SET collection='reference',extraction_required=0 WHERE id=?",(doc['id'],))
 db.execute('DELETE FROM schema_versions WHERE version=3')
 upgrade_arrival_notice_scope()
 after=bundle(cid)
 assert after['documents'][0]['id']==doc['id'] and after['documents'][0]['extraction_required']
 assert not after['case']['reviewed_at'] and not after['case']['requirements']['Booking']['required']
 assert after['case']['requirements']['Arrival Notice']['required']
 revision=after['case']['revision'];upgrade_arrival_notice_scope()
 assert bundle(cid)['case']['revision']==revision
