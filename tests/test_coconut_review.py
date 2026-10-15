import pytest
from backend import ai,extraction,pdf_bundle,consistency
from backend.normalize import normalize,decimal_text
from backend.validation import comparison_values

@pytest.mark.parametrize('raw,expected',[('FEB 11th 2020','2020-02-11'),('Jan, 31st 2020','2020-01-31'),('18th February 2020','2020-02-18')])
def test_named_ordinal_dates(raw,expected):
    value=normalize(raw,'date')
    assert value['text']==expected and not value['warning']
    assert normalize('02/03/2020','date')['warning']
    assert normalize('FEB 31st 2020','date')['warning']
    assert decimal_text('19,560') is None  # No global separator guess.

def test_contract_attachment_list_is_not_a_second_document():
    pages=[{'number':1,'text':'CONTRACT\nNo: TPV 03/EX'},
           {'number':2,'text':'ARTICLE 2 : Term of payment\nBy L/C\nARTICLE 4: Full set of documents required:\n- Bill of Lading\n- Commercial Invoice\n- Packing List\n- Certificate of Origin'}]
    assert pdf_bundle.suggest(pages)==[{'start':1,'end':2,'type':'Contract'}]
    assert not pdf_bundle.is_mixed(pages)
    assert pdf_bundle.is_mixed([pages[0],{'number':2,'text':'COMMERCIAL INVOICE\nInvoice No: INV-1'}])

def test_contract_reference_can_contain_source_printed_spaces():
    page={'number':1,'text':'CONTRACT\nNo: TPV 03/EX'}
    model=ai.DocumentResult.model_validate({'document_type':'Contract','classification_evidence':'CONTRACT','fields':[{'field_id':'contract_number','raw_value':'TPV 03/EX','source_page':1,'evidence_text':'No: TPV 03/EX'}],'items':[],'containers':[]})
    rows,rejected=ai.validated_subset(ai.ground_source_format(model,[page]),[page])
    assert not rejected and rows[0]['value']['text']=='TPV 03/EX'

def test_order_bank_consignee_is_never_promoted_to_buyer():
    page={'number':1,'text':'COMMERCIAL INVOICE\nInvoice No: INV-10\nCONSIGNEE: TO THE ORDER OF EXAMPLE BANK'}
    model=ai.DocumentResult.model_validate({'document_type':'Invoice','classification_evidence':'COMMERCIAL INVOICE','fields':[{'field_id':'buyer','raw_value':'TO THE ORDER OF EXAMPLE BANK','source_page':1,'evidence_text':'CONSIGNEE: TO THE ORDER OF EXAMPLE BANK'}],'items':[],'containers':[]})
    with pytest.raises(ValueError):ai.validate_result(model,[page])
    grounded=ai.ground_source_format(model,[page]);rows,rejected=ai.validated_subset(grounded,[page])
    assert not rows and not rejected

def test_contract_buyer_can_supply_shared_working_field_without_filling_invoice(monkeypatch):
    import json
    from backend.dictionary import default_requirements
    from backend.validation import build_checks
    monkeypatch.setattr(consistency,'ai_checks',lambda *args:[])
    docs=[{'id':'inv','type':'Invoice','name':'inv.pdf','status':'read','type_verified':0,'classification_method':'OpenAI'}, {'id':'contract','type':'Contract','name':'contract.pdf','status':'read','type_verified':0,'classification_method':'OpenAI'}]
    record={'id':'r','document_id':'contract','page':1,'kind':'field','key':'buyer','status':'auto_accepted','value':normalize('REAL BUYER')}
    checks=build_checks({'id':'case','mode':'FCL','requirements':json.dumps(default_requirements())},docs,[record])
    assert not any(c['key']=='missing:buyer' for c in checks)
    assert record['document_id']=='contract' and len([r for r in [record] if r['document_id']=='inv'])==0

def rows(key,values):
    return [{'id':str(i),'document_id':str(i),'page':1,'kind':'field','key':key,'status':'pending','value':normalize(v)} for i,v in enumerate(values)]

def test_port_formatting_does_not_merge_different_countries_or_ports():
    group=rows('port_loading',['CAT LAI, VIETNAM','CAT LAI PORT,VIET NAM','CAT LAI'])
    assert len(set(comparison_values('port_loading',group,group)))==1
    for other in ['CAT LAI, THAILAND','HO CHI MINH CITY, VIETNAM','ANY PORT OF HOCHIMINH CITY, VIETNAM']:
        different=group+[rows('port_loading',[other])[0]]
        assert len(set(comparison_values('port_loading',different,different)))>1
    mixed=group+[{**group[0],'id':'contract','value':normalize('ANY PORT OF HOCHIMINH CITY, VIETNAM')}]
    values=comparison_values('port_loading',mixed,mixed)
    assert len(set(values[:3]))==1 and values[3]!=values[0]

def test_vessel_suffix_needs_an_explicit_independent_voyage():
    group=rows('vessel',['AMALIA C','AMALIA C 074B'])
    assert len(set(comparison_values('vessel',group,group)))==2
    voyage={**group[0],'id':'voyage','key':'voyage','value':normalize('074B')}
    assert len(set(comparison_values('vessel',group,group+[voyage])))==1
    assert group[1]['value']['text']=='AMALIA C 074B'
    changed=rows('vessel',['AMALIA C','AMALIA C 075B'])
    assert len(set(comparison_values('vessel',changed,changed+[voyage])))==2
    check={'detail':{'sources':[{'id':r['id']} for r in group]}}
    assert not consistency.semantic_difference(check,group+[voyage])

def result(kind,fields=None,items=None):
    return ai.DocumentResult.model_validate({'document_type':kind,'classification_evidence':{'C/O':'CERTIFICATE OF ORIGIN','Booking':'BOOKING NOTE','Invoice':'COMMERCIAL INVOICE'}[kind], 'fields':fields or [],'items':items or [],'containers':[]})

def test_empty_slots_and_equipment_are_not_actionable_fields():
    page={'number':1,'text':'BOOKING NOTE\nCảng D/C : TAICHUNG\nSố lượng container : 1X20DC'}
    fields=[{'field_id':k,'raw_value':v,'source_page':1,'evidence_text':e} for k,v,e in [('eta','',''),('packages','1X20DC','Số lượng container : 1X20DC'),('port_discharge','TAICHUNG','Cảng D/C : TAICHUNG')]]
    grounded=ai.ground_source_format(result('Booking',fields),[page]);records,rejected=ai.validated_subset(grounded,[page])
    assert not rejected and [r['key'] for r in records]==['port_discharge']

def test_origin_declaration_number_and_certification_date_not_invoice_date():
    text='CERTIFICATE OF ORIGIN\n2. Consignee\nPURESUN\nNUMBER : 200218-03/02\n3. Means of transport\nNUMBER OF INVOICE\nNO: TPV03-02/EX\nDATE: FEB 11th 2020\nTRUONG PHU VINH here by certifies that above mention goods are of\nVietNamese origin.\n\nBENTRE VIETNAM FEB 18th 2020'
    fields=[{'field_id':'co_number','raw_value':'200218-03/02','source_page':1,'evidence_text':'NUMBER : 200218-03/02'}, {'field_id':'co_date','raw_value':'FEB 11th 2020','source_page':1,'evidence_text':'DATE: FEB 11th 2020'}]
    page={'number':1,'text':text}
    grounded=ai.ground_source_format(result('C/O',fields),[page]);records,rejected=ai.validated_subset(grounded,[page])
    assert not rejected
    assert next(r for r in records if r['key']=='co_date')['value']['text']=='2020-02-18'
    bad=result('C/O',[{**fields[0],'raw_value':'TPV03-02/EX','evidence_text':'NO: TPV03-02/EX'}])
    with pytest.raises(ValueError):ai.validate_result(bad,[page])

def test_joined_quantity_and_prefixed_currency_rebind_only_literal_item_cells():
    page={'number':1,'text':'COMMERCIAL INVOICE\nQUANTITY UNIT PRICE AMOUNT\nCOCONUT SAP 16,200KG USD 0.60/KG USD 9,720'}
    item={'line_number':'1','description':'COCONUT SAP','quantity':'16,200','unit':'KG','unit_price':'0.60','line_amount':'9,720','currency':'USD','source_page':1,'evidence_text':page['text'].splitlines()[2], 'cell_evidence':[{'attribute':'unit','evidence_text':'UNIT'},{'attribute':'currency','evidence_text':'AMOUNT'}]}
    grounded=ai.ground_source_format(result('Invoice',items=[item]),[page]);records,rejected=ai.validated_subset(grounded,[page])
    assert not rejected and any(r['kind']=='item' for r in records)
    # A price-unit suffix is not a quantity unit when the quantity has none.
    no_unit={**page,'text':page['text'].replace('16,200KG','16,200')}
    item['evidence_text']=no_unit['text'].splitlines()[2]
    grounded=ai.ground_source_format(result('Invoice',items=[item]),[no_unit])
    with pytest.raises(ValueError):ai.validate_result(grounded,[no_unit])

def test_packing_commodity_comes_from_unique_quantity_bound_cell():
    page={'number':1,'text':'PACKING LIST\nCOCONUT SAP AND EXTRACT\n(CLASS A)\nORIGIN OF VIETNAM\n120 DRUMS',
          'tables':[[['DESCRIPTION OF GOODS','NUMBER OF CASKS'],['COCONUT SAP AND EXTRACT\n(CLASS A)\nORIGIN OF VIETNAM','120 DRUMS']]]}
    item={'line_number':'1','description':'(CLASS A) ORIGIN OF VIETNAM','quantity':'120','unit':'DRUMS','source_page':1,'evidence_text':page['text']}
    model=ai.DocumentResult.model_validate({'document_type':'Packing List','classification_evidence':'PACKING LIST','fields':[],'items':[item],'containers':[]})
    grounded=ai.ground_source_format(model,[page]);records,rejected=ai.validated_subset(grounded,[page])
    assert not rejected and records[0]['value']['description']=='COCONUT SAP AND EXTRACT\n(CLASS A)'
    ambiguous={**page,'tables':[[*page['tables'][0],['ANOTHER COMMODITY','120 DRUMS']]],'text':page['text']+'\nANOTHER COMMODITY'}
    grounded=ai.ground_source_format(model,[ambiguous])
    assert grounded.items[0].description==item['description']  # Do not choose a row by count alone.


def test_merged_total_weight_header_keeps_per_drum_weight_separate():
    table=[['DESCRIPTION OF GOODS','NUMBER OF CASKS','',None,'TOTAL\nWEIGHT (KGS)',None],
           [None,None,'NET\nKGS','GROSS\nKGS','NET\nKGS','GROSS\nKGS'],
           ['COCONUT SAP','120 DRUMS','135','163','16,200','19,560'],
           ['REMARKS','* PACKING : 120 DRUMS\nTOTAL',None,None,None,None]]
    page={'number':1,'text':'PACKING LIST\n'+ '\n'.join(str(c) for row in table for c in row if c),'tables':[table]}
    model=ai.DocumentResult.model_validate({'document_type':'Packing List','classification_evidence':'PACKING LIST','fields':[], 'items':[], 'containers':[]})
    grounded=ai.ground_source_format(model,[page]);records,rejected=ai.validated_subset(grounded,[page])
    assert not rejected
    raw={r['key']:r['raw'] for r in records}
    assert raw=={'net_weight':'16,200','gross_weight':'19,560','packages':'120'}
    bad=ai.DocumentResult.model_validate({**model.model_dump(),'fields':[{'field_id':'net_weight','raw_value':'135','unit':'KGS','unit_evidence':'KGS','source_page':1,'evidence_text':'NET KGS 135'}]})
    with pytest.raises(ValueError):ai.validate_result(bad,[page])
    multi={**page,'tables':[[*table[:3],['OTHER GOODS','20 DRUMS','100','110','2,000','2,200'],table[-1]]]}
    assert not ai.packing_summary_fields([multi]).get('net_weight')  # Do not sum separate rows.


def test_readable_cells_rebind_interleaved_quotes_without_rewriting_values():
    cell='COCONUT SAP\nAND EXTRACT\n( CLASS A)'
    address='1. Consignor\nREAL EXPORTER LTD\n348D, STREET NAME\nVILLAGE, VIETNAM'
    page={'number':1,'text':'CONTRACT\nCOCONUT SAP 100 200\nAND EXTRACT\n( CLASS A)\n[PDF content-order layer]\n'+cell+'\n'+address,
          'tables':[[['Commodity','Quantity'],[cell,'100']],[[address,'CERTIFICATE OF ORIGIN']]]}
    fields=[{'field_id':'cargo_description','raw_value':'COCONUT SAP AND EXTRACT ( CLASS A)','source_page':1,'evidence_text':'COCONUT SAP 100 200\nAND EXTRACT\n( CLASS A)'}]
    model=ai.DocumentResult.model_validate({'document_type':'Contract','classification_evidence':'CONTRACT','fields':fields,'items':[],'containers':[]})
    fixed=ai.ground_source_format(model,[page]);records,rejected=ai.validated_subset(fixed,[page])
    assert not rejected and records[0]['raw']==fields[0]['raw_value'] and fixed.fields[0].evidence_text==cell
    no_layer={**page,'text':page['text'].split('[PDF content-order layer]')[0]}
    assert ai.ground_source_format(model,[no_layer]).fields[0].evidence_text==fields[0]['evidence_text']


def test_origin_adjective_uses_explicit_country_and_notify_address_is_not_name():
    page={'number':1,'text':'CERTIFICATE OF ORIGIN\nORIGIN OF VIETNAM\nVietNamese origin.\nNotify Party: EXAMPLE LTD\n5F, No. 5-4, Alley 10, Lane 30, Tung An Road, Taiwan'}
    fields=[{'field_id':'origin','raw_value':'VietNamese','source_page':1,'evidence_text':'VietNamese origin.'}]
    fixed=ai.ground_source_format(result('C/O',fields),[page]);records,rejected=ai.validated_subset(fixed,[page])
    assert not rejected and records[0]['raw']=='VIETNAM'
    wrong=result('C/O',[{**fields[0],'raw_value':'Chinese'}]);assert ai.ground_source_format(wrong,[page]).fields[0].raw_value=='Chinese'
    fields=[{'field_id':'notify_party','raw_value':'EXAMPLE LTD','source_page':1,'evidence_text':'Notify Party: EXAMPLE LTD'}, {'field_id':'notify_party','raw_value':page['text'].splitlines()[-1],'source_page':1,'evidence_text':page['text'].splitlines()[-1]}]
    model=ai.DocumentResult.model_validate({'document_type':'Arrival Notice','classification_evidence':'Notify Party','fields':fields,'items':[],'containers':[]})
    fixed=ai.ground_source_format(model,[page]);assert len(fixed.fields)==1 and fixed.fields[0].raw_value=='EXAMPLE LTD'


def test_warehouse_count_uses_quantity_after_packing_unit_never_row_index():
    page={'number':1,'text':'YEU CAU NHAP KHO\nTổng số tờ khai: 1\nNo. Description Unit Of Measure Quantity Weight Volume\n1 CT 56 515.080 7.380'}
    model=ai.DocumentResult.model_validate({'document_type':'Warehouse Receipt','classification_evidence':'YEU CAU NHAP KHO','fields':[{'field_id':'packages','raw_value':'1','source_page':1,'evidence_text':'Tổng số tờ khai: 1'}],'items':[],'containers':[]})
    records,rejected=ai.validated_subset(ai.ground_source_format(model,[page]),[page])
    assert not rejected and records[0]['raw']=='56' and records[0]['value']['unit']=='CTN'
    many={**page,'text':page['text']+'\n2 CT 20 100.00 1.20'}
    assert ai.ground_source_format(model,[many]).fields[0].raw_value=='1'  # No hidden sum or preferred row.


def test_two_tier_weight_headers_and_joined_quantity_are_literal_sources():
    page={'number':1,'text':'PACKING LIST\nITEM NO.PACKAGES N.WEIGHT G.WEIGHT\n(KGS) (KGS)\nTOTAL 2112 24288 25200\n42240SETS',
          'tables':[[['ITEM','NO.PACKAGES','N.WEIGHT','G.WEIGHT'],[None,None,'(KGS)','(KGS)'],['TOTAL','2112','24288','25200']]]}
    model=ai.DocumentResult.model_validate({'document_type':'Packing List','classification_evidence':'PACKING LIST','fields':[{'field_id':'goods_quantity','raw_value':'42240 SETS','source_page':1,'evidence_text':'42240SETS'}],'items':[],'containers':[]})
    records,rejected=ai.validated_subset(ai.ground_source_format(model,[page]),[page])
    assert not rejected and {r['key']:r['raw'] for r in records}=={'goods_quantity':'42240','packages':'2112','net_weight':'24288','gross_weight':'25200'}


def test_unknown_classification_does_not_fabricate_a_document_or_block_other_files():
    page={'number':1,'text':'Warehouse screen\nStatus: SUCCESS'}
    model=ai.DocumentResult.model_validate({'document_type':'Other','classification_evidence':'Not actually on the page','fields':[],'items':[],'containers':[]})
    fixed=ai.ground_source_format(model,[page]);assert fixed.document_type=='Other' and fixed.classification_evidence=='Warehouse screen'
    assert ai.validate_result(fixed,[page])==[]
    assert extraction.classify([{'number':1,'text':'PURCHASE CONTRACT\nNo: QL/DS2602'}])[0]=='Contract'
    assert extraction.classify([{'number':1,'text':'WAYBILL\nNON NEGOTIABLE\nPORT OF LOADING\nTRANSPORT BILL OF LADING'}])[0]=='Sea Waybill'


def test_excel_empty_formatted_tail_is_not_data_but_distant_values_are_checked(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    from backend import reader
    wb=Workbook();s=wb.active;s['A1']='Commercial Invoice';s['B2']='INV-1';s['IV449'].fill=PatternFill('solid',fgColor='FFFFFF')
    p=tmp_path/'formatted.xlsx';wb.save(p)
    pages=reader.read_document(p);assert len(pages[0]['tables'][0])==2 and len(pages[0]['tables'][0][0])==2
    s['IV449']='real data';wb.save(p);wb.close()
    with pytest.raises(ValueError,match='cột'):reader.read_document(p)
