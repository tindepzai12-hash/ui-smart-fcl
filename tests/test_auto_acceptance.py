import io,json
from openpyxl import load_workbook
from backend import acceptance,ai,db,extraction,validation,consistency
from backend.app import bundle,revalidate
from backend.normalize import normalize
from test_core import client,H,new,upload,process
from fixtures import INVOICE,PACKING,BL

def test_single_source_ai_goes_directly_to_workbook_without_manual_review(monkeypatch):
    cid=client.post('/api/cases',headers=H,json={'name':'Auto accept source'}).json()['case']['id']
    upload(cid,[('invoice.txt',INVOICE.encode())])
    monkeypatch.setattr(ai,'key',lambda:'mock')
    def analyze(case,doc,pages):
        result=ai.DocumentResult.model_validate({'document_type':'Invoice','classification_evidence':'INVOICE','fields':[],'items':[],'containers':[]})
        rows=extraction.extract(pages,'Invoice')
        for r in rows:r['method']='OpenAI'
        return result,rows,{'source':'test','rejected':[]}
    monkeypatch.setattr(ai,'analyze',analyze)
    monkeypatch.setattr(consistency,'run_ai_review',lambda *args:'')
    b,j=process(cid,'openai');assert not j['errors']
    row=next(r for r in b['records'] if r['key']=='invoice_number')
    assert row['status']=='auto_accepted' and row['evidence']
    assert b['documents'][0]['classification_method']=='OpenAI'
    assert not any(c['key'].startswith('classification:') for c in b['checks'])
    assert not any('bản ghi chưa' in s for s in b['blockers'])
    response=client.post('/api/cases/'+cid+'/export',headers=H,json={'mode':'Draft'});assert response.status_code==200
    wb=load_workbook(io.BytesIO(client.get(response.json()['download']).content))
    assert any(r[0]=='invoice_number' and r[5]=='auto_accepted' for r in wb.worksheets[1].iter_rows(values_only=True));wb.close()
    # Reading again replaces machine-accepted rows, never duplicates them.
    again,_=process(cid,'openai')
    assert len(again['records'])==len(b['records'])

def test_checkbox_saves_without_note_but_never_approves_missing_fields():
    cid=new();upload(cid,[('inv.txt',INVOICE.encode()),('pl.txt',PACKING.encode()),('bl.txt',BL.replace('1 MT','1100 KG').encode())]);b,_=process(cid)
    check=next(c for c in b['checks'] if c['key']=='field:gross_weight')
    ids={s['id'] for s in check['detail']['sources']};versions={r['id']:r['version'] for r in b['records'] if r['id'] in ids}
    url='/api/checks/'+check['id']+'/review'
    assert client.post(url,headers=H,json={'versions':versions}).status_code==422
    r=client.post(url,headers=H,json={'versions':versions,'confirmed':True});assert r.status_code==200,r.text
    saved=next(c for c in r.json()['checks'] if c['key']==check['key']);assert saved['result']=='Mismatch' and saved['resolution']
    missing=next((c for c in r.json()['checks'] if c['result']=='Missing'),None)
    if missing:assert client.post('/api/checks/'+missing['id']+'/review',headers=H,json={'versions':{},'confirmed':True}).status_code==422
    row=next(r for r in b['records'] if r['key']=='invoice_number')
    assert client.patch('/api/records/'+row['id'],headers=H,json={'value':'','version':row['version']}).status_code==422

def test_unreadable_dates_units_and_empty_required_values_cannot_auto_pass():
    for key,raw,kind in [('booking','','text'),('gross_weight','abc','weight'),('gross_weight','1000','weight'),('eta','31/02/2020','date')]:
        assert not acceptance.readable({'kind':'field','key':key,'value':normalize(raw,kind)})
    assert not acceptance.readable({'kind':'item','value':{'description':'Cargo','quantity':'120','unit':''}})
    for key,raw in [('cargo_weight','515.080'),('volume','7.380')]:
        value=normalize(raw,'weight' if key=='cargo_weight' else 'quantity')
        assert acceptance.readable({'kind':'field','key':key,'value':value}) and value['unit'] is None and value['text']==raw
    assert acceptance.readable({'kind':'field','key':'gross_weight','value':normalize('19,560 KGS','weight')})
    assert normalize('19,560 KGS','weight')['number'] is None  # Accepted as literal text, no guessed numeric conversion.
    date=normalize('02/03/2020','date')
    assert acceptance.readable({'kind':'field','key':'eta','value':date}) and date['text']=='02/03/2020'

def test_equal_literal_figures_match_and_single_source_does_not_become_an_exception():
    rows=[{'id':str(i),'document_id':str(i),'page':1,'kind':'field','key':'gross_weight','status':'pending','value':normalize('19,560 KGS','weight')} for i in range(2)]
    assert len(set(validation.comparison_values('gross_weight',rows,rows)))==1
    check={'detail':{'sources':[{'id':r['id']} for r in rows]}}
    assert not consistency.semantic_difference(check,rows)
    rows[1]['value']=normalize('19,660 KGS','weight')
    assert len(set(validation.comparison_values('gross_weight',rows,rows)))==2
    assert consistency.semantic_difference(check,rows)

def test_missing_required_manual_entry_is_atomic_and_never_accepts_blank():
    cid=client.post('/api/cases',headers=H,json={'name':'Manual booking'}).json()['case']['id']
    upload(cid,[('booking.txt',b'BOOKING CONFIRMATION\nBooking No: UNKNOWN\nVessel: EXAMPLE')]);b,_=process(cid)
    did=b['documents'][0]['id']
    body={'kind':'field','key':'booking','value':'','note':'Nhân viên đã kiểm tra bản gốc.','document_id':did,'page':1,'confirmed':True}
    before=len(b['records'])
    assert client.post('/api/cases/'+cid+'/records',headers=H,json=body).status_code==422
    assert len(bundle(cid)['records'])==before
    body['value']='REAL-BOOKING-100'
    res=client.post('/api/cases/'+cid+'/records',headers=H,json=body);assert res.status_code==200,res.text
    assert any(r['status']=='verified' and r['value']['text']==body['value'] for r in res.json()['records'])
    body.update(key='gross_weight',value='unreadable KG')
    assert client.post('/api/cases/'+cid+'/records',headers=H,json=body).status_code==422

def test_restore_endpoint_refuses_to_replace_existing_work():
    cid=new()
    before=client.get('/api/cases').json()
    res=client.post('/api/backups/restore-empty',headers=H,files={'file':('backup.zip',b'invalid','application/zip')})
    assert res.status_code==409
    assert client.get('/api/cases').json()==before

def test_restore_empty_server_recovers_backup_and_rejects_bad_archive(monkeypatch):
    cid=new()
    export=client.post('/api/backup',headers=H).json()
    backup=client.get(export['download']).content
    import tempfile
    from pathlib import Path
    destination=Path(tempfile.mkdtemp(prefix='restore-api-',dir=db.ROOT/'tmp'))
    monkeypatch.setattr(db,'DATA',destination)
    for folder in ['files','exports','templates','backups']:(destination/folder).mkdir()
    db.migrate()
    bad=client.post('/api/backups/restore-empty',headers=H,files={'file':('bad.zip',b'bad','application/zip')})
    assert bad.status_code==422 and not client.get('/api/cases').json()
    res=client.post('/api/backups/restore-empty',headers=H,files={'file':('backup.zip',backup,'application/zip')})
    assert res.status_code==200,res.text
    assert any(c['id']==cid for c in client.get('/api/cases').json())
