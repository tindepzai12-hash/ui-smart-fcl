import io, json, zipfile
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend.app import app, bundle
from backend import db, extraction, validation
from backend.normalize import normalize
from fixtures import generate, INVOICE, PACKING, BL
from test_core import new, upload, process, H, client
from scripts.restore import restore

def test_template_mapping_and_foreign_scope(tmp_path):
    generate(tmp_path)
    id = new()
    upload(id, [('invoice.txt', INVOICE.encode())])
    b, j = process(id)
    r = next(r for r in b['records'] if r['key'] == 'invoice_number')
    t = client.post('/api/templates', headers=H, files={'file': ('internal-template.xlsx', (tmp_path / 'internal-template.xlsx').read_bytes())})
    assert t.status_code == 200
    mapping = [{'sheet': 'Internal mapping', 'cell': 'B3', 'record_id': r['id']}]
    exported = client.post('/api/templates/' + t.json()['id'] + '/export', headers=H, json={'case_id': id, 'mapping': mapping})
    assert exported.status_code == 200, exported.text
    wb = load_workbook(io.BytesIO(client.get(exported.json()['download']).content))
    assert wb['Internal mapping']['B3'].value == '000123-A'
    assert 'UI_SMART_NGUON' in wb.sheetnames
    # Ensure trace sheet records human-readable source document name instead of raw UUID only
    trace_values = [cell for row in wb['UI_SMART_NGUON'].iter_rows(min_row=2, values_only=True) for cell in row]
    assert 'invoice.txt' in trace_values
    wb.close()
    assert client.post('/api/templates/' + t.json()['id'] + '/export', headers=H, json={'case_id': new(), 'mapping': mapping}).status_code == 422
    assert client.post('/api/templates/' + t.json()['id'] + '/export', headers=H, json={'case_id': id, 'mapping': mapping * 2}).status_code == 422

def test_backup_restore_and_traversal(tmp_path):
    archive = client.post('/api/backup', headers=H).json()
    p = tmp_path / 'backup.zip'
    p.write_bytes(client.get(archive['download']).content)
    destination = db.ROOT / 'tmp' / 'restore-test'
    old = restore(p, destination, check_server=False)
    import sqlite3
    with sqlite3.connect(destination / 'smart.sqlite') as c:
        assert c.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
        assert c.execute('SELECT count(*) FROM cases').fetchone()[0] > 0
        for (f,) in c.execute('SELECT path FROM documents'):
            assert Path(f).is_file()
            assert Path(f).is_relative_to(destination)
    evil = tmp_path / 'evil.zip'
    with zipfile.ZipFile(evil, 'w') as z:
        z.writestr('../outside.txt', 'escape')
    with pytest.raises(ValueError):
        restore(evil, db.ROOT / 'tmp' / 'evil-restore', check_server=False)

def test_resolution_invalidates_on_edit():
    id = new()
    upload(id, [('invoice.txt', INVOICE.encode()), ('packing.txt', PACKING.encode()), ('bl.txt', BL.replace('1 MT', '1100 KG').encode())])
    b, _ = process(id)
    check = next(c for c in b['checks'] if c['key'] == 'field:gross_weight')
    assert check['result'] == 'Mismatch'
    assert client.post('/api/checks/' + check['id'] + '/resolve', headers=H, json={'note': 'Căn cứ kiểm thử: đang chờ sửa B/L'}).status_code == 200
    r = next(r for r in bundle(id)['records'] if r['key'] == 'gross_weight' and r['value']['number'] == '1100')
    assert client.patch('/api/records/' + r['id'], headers=H, json={'value': '1200 KG', 'note': 'Theo bản B/L thay thế', 'version': r['version']}).status_code == 200
    check = next(c for c in bundle(id)['checks'] if c['key'] == 'field:gross_weight')
    assert not check['resolution']

def test_quantities_units_roles_not_false_conflict():
    id = new()
    upload(id, [('invoice.txt', (INVOICE + '\nQuantity: 1000 PCS').encode()), ('packing.txt', (PACKING + '\nQuantity: 1000 PCS').encode()), ('bl.txt', BL.encode())])
    b, _ = process(id)
    checks = {c['key']: c for c in b['checks']}
    assert checks['field:packages']['result'] == 'Matched'
    assert checks['field:goods_quantity']['result'] == 'Matched'
    assert not any(c['key'] in ['field:seller', 'field:shipper'] and c['result'] == 'Mismatch' for c in b['checks'])

def test_manual_empty_and_evidence_gate():
    id = new()
    assert client.post('/api/cases/' + id + '/records', headers=H, json={'kind': 'field', 'key': 'invoice_number', 'value': '', 'note': 'reason'}).status_code == 422
    res = client.post('/api/cases/' + id + '/records', headers=H, json={'kind': 'item', 'key': 'goods_item', 'value': {'description': 'Manual item'}, 'note': 'Theo nguồn mẫu'})
    assert res.status_code == 200
    r = res.json()['records'][0]
    assert client.post('/api/records/' + r['id'] + '/verify', headers=H, json={'status': 'verified', 'version': r['version']}).status_code == 422

def test_secret_not_echoed_and_explicit_ai_consent():
    from backend import ai
    secret = 'synthetic-private-key-' + ('x' * 510)
    res = client.post('/api/ai/configure', headers=H, json={'api_key': secret, 'model': 'gpt-4.1-mini'})
    assert res.status_code == 422
    assert 'synthetic-private-key' not in res.text
    id = new()
    upload(id, [('invoice.txt', INVOICE.encode())])
    ai.configure('synthetic-consent-key', 'gpt-4.1-mini')
    try:
        res = client.post('/api/cases/' + id + '/process', headers=H, json={'mode': 'openai'})
        assert res.status_code == 422
        assert 'đồng ý' in res.text
    finally:
        ai.remove()

def test_checklist_exemption_and_required_field_distinction():
    id = new()
    upload(id, [('invoice.txt', INVOICE.encode()), ('packing.txt', PACKING.encode())])
    b, _ = process(id)
    checks = {c['key']: c for c in b['checks']}
    assert checks['document:B/L']['result'] == 'Missing'
    assert checks['missing:bl_number']['result'] == 'Missing'
    # Missing check cannot be bypassed by resolve note while document is still required
    assert client.post('/api/checks/' + checks['missing:bl_number']['id'] + '/resolve', headers=H, json={'note': 'Cố gắng bỏ qua khi chưa miễn B/L'}).status_code == 422
    # Exempt B/L in checklist with reason -> document:B/L becomes Not Applicable, and missing:bl_number becomes Needs Verification
    req_res = client.patch('/api/cases/' + id + '/requirements/B%2FL', headers=H, json={'required': False, 'note': 'Hồ sơ kiểm tra trước khi phát hành B/L'})
    assert req_res.status_code == 200
    checks_after = {c['key']: c for c in req_res.json()['checks']}
    assert checks_after['document:B/L']['result'] == 'Not Applicable'
    assert checks_after['missing:bl_number']['result'] == 'Needs Verification'
    # Staff can now explicitly confirm/resolve the exempted bl_number field with a reason
    res_resolve = client.post('/api/checks/' + checks_after['missing:bl_number']['id'] + '/resolve', headers=H, json={'note': 'Xác nhận miễn số B/L ở bước kiểm tra thương mại'})
    assert res_resolve.status_code == 200

def test_ambiguous_slash_date_and_resolution_guards():
    # Ambiguous slash date (01/10/2026) keeps parsed text 2026-10-01 AND sets warning for human verification
    ambig = normalize('01/10/2026', 'date')
    assert ambig['text'] == '2026-10-01'
    assert ambig['warning'] is not None
    # Unambiguous slash date (25/10/2026 where day > 12) and ISO date have no warning
    assert normalize('25/10/2026', 'date')['warning'] is None
    assert normalize('2026-10-01', 'date')['warning'] is None

    # Classification checks cannot be resolved by a note instead of confirming document classification
    id = new()
    upload(id, [('invoice.txt', INVOICE.encode()), ('packing.txt', PACKING.encode()), ('bl.txt', BL.encode())])
    b, _ = process(id)
    cls_check = next(c for c in b['checks'] if c['key'].startswith('classification:'))
    assert client.post('/api/checks/' + cls_check['id'] + '/resolve', headers=H, json={'note': 'Bỏ qua xác nhận loại'}).status_code == 400

    # Export stats include physical_containers=2 vs containers=4
    exp = client.post('/api/cases/' + id + '/export', headers=H, json={'mode': 'Draft'}).json()
    assert exp['stats']['physical_containers'] == 2
    assert exp['stats']['containers'] == 4


def test_reading_and_whitespace_cannot_bypass_review():
    id = new()
    upload(id, [('invoice.txt', INVOICE.encode())])
    b = bundle(id)
    check = next(c for c in b['checks'] if c['key'].startswith('reading:'))
    result = client.post('/api/checks/' + check['id'] + '/resolve', headers=H, json={'note': 'Không đọc chứng từ'})
    assert result.status_code == 400
    assert 'ghi chú' in result.json()['detail']
    assert client.patch('/api/cases/' + id + '/requirements/B%2FL', headers=H, json={'required': False, 'note': '     '}).status_code == 422
    assert client.post('/api/checks/' + check['id'] + '/resolve', headers=H, json={'note': '       '}).status_code == 422
    assert not any(s.startswith(('document:', 'reading:', 'classification:', 'missing:')) for s in b['blockers'])


def test_valid_bl_waiver_can_reach_reviewed_but_not_waive_other_dependencies():
    id = new()
    upload(id, [('invoice.txt', INVOICE.encode()), ('packing.txt', PACKING.encode())])
    b, _ = process(id)
    for d in b['documents']:
        res = client.patch('/api/documents/' + d['id'] + '/classification', headers=H, json={'type': d['type']})
        assert res.status_code == 200
    res = client.patch('/api/cases/' + id + '/requirements/B%2FL', headers=H, json={'required': False, 'note': 'Bản làm việc chưa nhận B/L; kiểm tra theo hồ sơ thương mại'})
    assert res.status_code == 200
    b = res.json()
    assert next(c for c in b['checks'] if c['key'] == 'missing:bl_number')['result'] == 'Needs Verification'
    for r in b['records']:
        assert client.post('/api/records/' + r['id'] + '/verify', headers=H, json={'status': 'verified', 'version': r['version'], 'note': 'Đối chiếu từng giá trị với chứng từ mẫu'}).status_code == 200
    for c in bundle(id)['checks']:
        if c['result'] in ['Mismatch', 'Needs Verification'] and not c['resolution']:
            assert client.post('/api/checks/' + c['id'] + '/resolve', headers=H, json={'note': 'Đã kiểm tra bộ mẫu; ghi nhận phạm vi và lý do miễn theo hồ sơ'}).status_code == 200
    reviewed = client.post('/api/cases/' + id + '/review', headers=H)
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()['case']['status'] == 'Reviewed'
    # Removing invoice evidence must still require invoice fields, despite a B/L waiver.
    checks = validation.build_checks({'mode': 'FCL', 'requirements': {'Invoice': {'required': True, 'note': 'Cần hóa đơn'}, 'B/L': {'required': False, 'note': 'Có lý do miễn'}}}, [], [])
    assert next(c for c in checks if c['key'] == 'missing:invoice_date')['result'] == 'Missing'


def test_changing_document_type_reopens_protected_records():
    id = new()
    upload(id, [('invoice.txt', INVOICE.encode())])
    b, _ = process(id)
    r = next(r for r in b['records'] if r['key'] == 'invoice_number')
    assert client.post('/api/records/' + r['id'] + '/verify', headers=H, json={'status': 'verified', 'version': r['version']}).status_code == 200
    changed = client.patch('/api/documents/' + b['documents'][0]['id'] + '/classification', headers=H, json={'type': 'Other'}).json()
    preserved = next(x for x in changed['records'] if x['id'] == r['id'])
    assert preserved['status'] == 'pending'
    assert preserved['version'] > r['version']
    assert preserved['value'] == r['value']
    assert not changed['case']['reviewed_at']


def test_normalized_physical_containers_weights_and_safe_samples():
    case = {'mode': 'FCL', 'requirements': {}}
    containers = [
        {'id': 'r1', 'kind': 'container', 'status': 'pending', 'document_id': 'd1', 'page': 1, 'value': {'container_number': 'mscu 1234567', 'seal_number': 'S1', 'gross_weight': '1000 KG'}},
        {'id': 'r2', 'kind': 'container', 'status': 'pending', 'document_id': 'd2', 'page': 1, 'value': {'container_number': 'MSCU1234567', 'seal_number': 'S1', 'gross_weight': '1 MT'}},
    ]
    checks = validation.build_checks(case, [], containers)
    assert len([c for c in checks if c['key'].startswith('container:')]) == 1
    assert next(c for c in checks if c['key'] == 'container-weight:MSCU1234567')['result'] == 'Matched'
    containers[1]['value']['gross_weight'] = '1100 KG'
    checks = validation.build_checks(case, [], containers)
    assert next(c for c in checks if c['key'] == 'container-weight:MSCU1234567')['result'] == 'Mismatch'
    containers[1]['value']['gross_weight'] = '1100'
    assert next(c for c in validation.build_checks(case, [], containers) if c['key'] == 'container-weight:MSCU1234567')['result'] == 'Needs Verification'
    assert client.get('/api/samples/invoice.txt').status_code == 200
    assert 'SYNTHETIC' in client.get('/api/samples/invoice.txt').text
    assert client.get('/api/samples/malicious.txt').status_code == 404
    assert client.get('/api/samples/..%2Fbackend%2Fapp.py').status_code == 404
