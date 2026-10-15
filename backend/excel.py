"""Internal workbook engine. Uses openpyxl explicitly requested in master prompt."""
import json
from decimal import Decimal
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from .dictionary import FIELDS, LEGAL, REVIEW_FIELDS, commercial_profile, supports
from .validation import human_check_label, container_number
from .db import now

SHEETS = [
    'TONG_QUAN_HO_SO',
    'THONG_TIN_KHAI_BAO',
    'CHI_TIET_HANG_HOA',
    'CONTAINER_VAN_CHUYEN',
    'DOI_CHIEU_CHUNG_TU',
    'THIEU_VA_CAN_XAC_MINH',
    'NGUON_VA_LICH_SU',
]

def safe(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        value = ''.join(c for c in value if ord(c) >= 32 or c in '\n\t')[:32000]
        # Keep codes/text as strings and neutralize spreadsheet formula prefixes.
        if value.lstrip().startswith(('=', '+', '-', '@')):
            value = "'" + value
    return value

def value_text(value):
    if not isinstance(value, dict):
        return str(value)
    if value.get('number') is not None:
        return value['number'] + (' ' + str(value['unit']) if value.get('unit') else '')
    if 'container_number' in value:
        parts = [value.get('container_number', '')]
        if value.get('seal_number'):
            parts.append('Seal: ' + str(value['seal_number']))
        if value.get('container_type'):
            parts.append('Loại: ' + str(value['container_type']))
        return ' · '.join(p for p in parts if p)
    if 'description' in value:
        return f"{value.get('description', '')} ({value.get('quantity', '')} {value.get('unit', '')})".strip()
    return value.get('text', '')

def format_sources(sources):
    if not sources:
        return '—'
    lines = []
    for s in sources:
        doc = s.get('document', 'Nhập tay')
        pg = f" (tr.{s['page']})" if s.get('page') else ''
        val = value_text(s.get('value', ''))
        lines.append(f"{doc}{pg}: {val}" if val else f"{doc}{pg}")
    return '\n'.join(lines)

def sheet(wb, name, headers, data):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in data:
        ws.append([safe(v) for v in row])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='155EEF')
        cell.font = Font(name='Calibri', color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if isinstance(cell.value, str):
                cell.number_format = '@'
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F1F5F9')
    for col in ws.columns:
        width = max(len(str(c.value or '')) for c in list(col)[:100])
        ws.column_dimensions[col[0].column_letter].width = min(60, max(18, width + 2))
    return ws

def numeric(raw, verified):
    if not verified or raw in [None, '']:
        return raw or ''
    try:
        d = Decimal(str(raw))
        # Excel has 15 significant digits. Preserve more precise data as text.
        if not d.is_finite() or len(d.as_tuple().digits) > 15:
            return str(raw)
        return float(d)
    except Exception:
        return str(raw)

def build(path, case, documents, records, checks, audit, mode='Draft', workbook_type='A', selected=None, version=1):
    requirements=json.loads(case['requirements']) if isinstance(case['requirements'],str) else case['requirements']
    commercial=commercial_profile(requirements,case['mode'])
    names=list(SHEETS)
    if commercial: names[1]='THONG_TIN_CHUNG_TU'
    expected=[n for i,n in enumerate(names) if not (commercial and i==3)]
    wb = Workbook()
    wb.remove(wb.active)
    docs = {d['id']: d for d in documents}
    active = [r for r in records if r['status'] != 'rejected' and (not commercial or (r['kind'] != 'container' and (r['kind'] != 'field' or r['key'] in REVIEW_FIELDS)))]
    fields = [r for r in active if r['kind'] == 'field' and (not selected or r['key'] in selected)]
    warnings = [c for c in checks if c['result'] in ['Mismatch', 'Missing', 'Needs Verification'] and not c.get('resolution')]
    resolved_conflicts = [c for c in checks if c['result'] in ['Mismatch', 'Needs Verification'] and c.get('resolution')]
    cont_records = [r for r in active if r['kind'] == 'container']
    unique_containers = sorted({container_number(r['value'].get('container_number')) for r in cont_records if container_number(r['value'].get('container_number'))})

    sheet(
        wb,
        names[0],
        ['Chỉ tiêu', 'Giá trị'],
        [
            ['Case ID', case['id']],
            ['Tên hồ sơ', case['name']],
            ['Khách hàng', case['customer']],
            ['Nghiệp vụ', case['mode']],
            ['Số chứng từ', len(documents)],
            *([] if commercial else [['Container vật lý (duy nhất)', f"{len(unique_containers)} container vật lý / {len(cont_records)} bản ghi nguồn"]]),
            *([['Checklist', 'Packing List · Booking · Invoice · Hợp đồng/PO · C/O (nếu có)']] if commercial else []),
            ['Phương pháp', 'Local / OpenAI theo từng nguồn'],
            ['Tạo lúc', case['created_at']],
            ['Duyệt lúc', case.get('reviewed_at') or 'Chưa hoàn thành duyệt Reviewed'],
            ['Xuất lúc', now()],
            ['Trạng thái xuất', mode],
            ['Loại workbook', workbook_type],
            ['Cảnh báo chưa xử lý', len(warnings)],
            ['Xung đột đã ghi nhận lý do xử lý', len(resolved_conflicts)],
            ['Phiên bản', version],
            ['Phạm vi', 'Hồ sơ làm việc nội bộ, không phải tờ khai chính thức hay mẫu import VNACCS.'],
            ['Dữ liệu thương mại', 'Chỉ chia sẻ workbook với người được phép.'],
            ['Ý nghĩa phiên bản', 'Bản Draft giữ nguyên các bản ghi chờ xác minh và cảnh báo. Bản Reviewed xác nhận nhân viên đã kiểm tra hết các cổng nội bộ.'],
        ],
    )
    grouped = {}
    for r in fields:
        grouped.setdefault(r['key'], []).append(r)
    data = []
    for key in (selected or (REVIEW_FIELDS if commercial else list(FIELDS))):
        if key not in FIELDS or (commercial and key not in REVIEW_FIELDS): continue
        group = grouped.get(key, [])
        if group:
            for r in group:
                data.append([
                    key,
                    FIELDS[key]['label'],
                    value_text(r['value']),
                    r['value'].get('unit'),
                    docs.get(r['document_id'], {}).get('name', 'Nhập tay'),
                    r['status'],
                    r['method'],
                    r['note'],
                ])
        else:
            blank_status = 'Missing'
            if commercial:
                missing_check = next((c for c in warnings if c['key'] == 'missing:' + key), None)
                # An absent extraction is not proof that a printed field is inapplicable.
                uncertain_documents = any(d['status'] != 'read' or not d['type_verified'] or d['type'] == 'Other' for d in documents)
                applicable_documents = any(supports(d['type'], kind) for d in documents for kind in FIELDS[key]['documents'])
                blank_status = missing_check['result'] if missing_check else ('Not Yet Checked' if uncertain_documents or applicable_documents or not documents else 'Not Applicable')
            data.append([
                key,
                FIELDS[key]['label'],
                '',
                '',
                '',
                blank_status,
                '',
                'Nhập/xác minh theo căn cứ chính thức' if key in LEGAL else (
                    'Không có loại chứng từ áp dụng trong bộ đã xác nhận' if commercial and blank_status == 'Not Applicable'
                    else 'Chưa ghi nhận; kiểm bản gốc trước khi kết luận không áp dụng' if commercial else 'Chưa trích xuất'),
            ])
    ws = sheet(wb, names[1], ['Field ID', 'Tên chỉ tiêu', 'Giá trị', 'Đơn vị', 'Nguồn', 'Trạng thái', 'Phương pháp', 'Ghi chú'], data)
    dv = DataValidation(type='list', formula1='"pending,verified,rejected,Missing,Needs Verification,Not Yet Checked,Not Applicable"')
    ws.add_data_validation(dv)
    dv.add(f'F2:F{max(2, ws.max_row)}')
    items = []
    for r in active:
        if r['kind'] != 'item':
            continue
        v = r['value']
        verified = r['status'] == 'verified'
        items.append(
            [
                v.get(k, '') if k not in ['quantity', 'unit_price', 'line_amount'] else numeric(v.get(k, ''), verified)
                for k in ['line_number', 'product_code', 'description', 'quantity', 'unit', 'unit_price', 'line_amount', 'currency', 'origin', 'hs_code']
            ]
            + [r['status'], docs.get(r['document_id'], {}).get('name', 'Nhập tay'), r['page'], v.get('scope', 'invoice')]
        )
    ws = sheet(
        wb,
        names[2],
        ['STT', 'Mã hàng', 'Mô tả', 'Số lượng', 'Đơn vị', 'Đơn giá', 'Thành tiền', 'Tiền tệ', 'Xuất xứ', 'Mã HS', 'Trạng thái', 'Nguồn', 'Trang', 'Phạm vi'],
        items,
    )
    for row in ws.iter_rows(min_row=2):
        for i in [3, 5, 6]:
            if isinstance(row[i].value, (int, float)):
                row[i].number_format = '#,##0.############'
    containers = []
    for r in active:
        if r['kind'] == 'container':
            v = r['value']
            containers.append(
                [v.get(k, '') for k in ['container_number', 'seal_number', 'container_type', 'packages', 'gross_weight', 'cbm', 'bl_number']]
                + [docs.get(r['document_id'], {}).get('name', 'Nhập tay'), r['status'], r['page']]
            )
    if not commercial: sheet(wb, names[3], ['Container', 'Seal', 'Loại container', 'Số kiện', 'Gross Weight', 'CBM', 'B/L', 'Nguồn', 'Trạng thái', 'Trang'], containers)
    sheet(
        wb,
        names[4],
        ['Quy tắc', 'Nguồn và giá trị', 'Kết quả', 'Mức độ', 'Lý do', 'Ghi chú xử lý'],
        [[f"{c['key']} ({human_check_label(c['key'], docs)})", format_sources(c['detail'].get('sources', [])), c['result'], c['severity'], c['detail']['reason'], c['resolution']] for c in checks],
    )
    missing = [[c['key'], 'Checklist nội bộ', c['detail']['reason'], format_sources(c['detail'].get('sources', [])) if c['detail'].get('sources') else 'Theo chứng từ', c['severity'], c['result'], c['resolution'] or 'Kiểm tra/bổ sung'] for c in warnings]
    for key in ([] if commercial else LEGAL):
        if not grouped.get(key):
            missing.append([key, 'Chuẩn bị khai báo', 'Không tự suy đoán mã/thuế/trị giá', 'Nguồn chính thức', 'Cần nhân viên nghiệp vụ', 'Missing', 'Tra cứu và xác minh; không phải mọi chỉ tiêu đều áp dụng'])
    sheet(wb, names[5], ['Thiếu/cần kiểm tra', 'Phục vụ', 'Lý do', 'Nguồn dự kiến', 'Ưu tiên', 'Trạng thái', 'Thao tác'], missing)
    history = [
        [
            r['key'],
            value_text(r['value']),
            docs.get(r['document_id'], {}).get('name', 'Nhập tay'),
            r['page'],
            r['evidence'],
            r['method'],
            r['status'],
            r['updated_at'],
            r['raw'] if isinstance(r['raw'], str) else value_text(r['raw']),
            r['note'],
        ]
        for r in records if not commercial or (r['kind']!='container' and (r['kind']!='field' or r['key'] in REVIEW_FIELDS))
    ]
    history += [[a['action'], a.get('after_json', ''), '', '', '', a['actor'], '', a['created_at'], a.get('before_json', ''), a['reason']] for a in audit]
    for doc in documents:
        if doc.get('source'):
            source = doc['source']
            history.append(['Nguồn PDF gốc', doc['name'], source['name'], f"{source['start_page']}-{source['end_page']}", 'SHA-256: ' + source['hash'], 'Tách PDF / giữ bản gốc', 'Chờ xác minh theo từng bản ghi', '', '', 'Khoảng trang trên PDF gốc; trang chứng từ con bắt đầu từ 1.'])
    sheet(
        wb,
        names[6],
        ['Trường/thao tác', 'Giá trị hiện tại/sau', 'Nguồn', 'Trang', 'Bằng chứng', 'Phương pháp/người sửa', 'Trạng thái', 'Thời gian', 'Giá trị gốc/trước', 'Ghi chú'],
        history,
    )
    wb.save(path)
    reopened = load_workbook(path)
    assert reopened.sheetnames == expected
    assert not any(c.data_type == 'f' for s in reopened for row in s for c in row)
    reopened.close()
    return {'sheets': expected, 'fields': len(fields), 'items': len(items), 'containers': len(containers), 'physical_containers': len(unique_containers), 'warnings': len(warnings)}
