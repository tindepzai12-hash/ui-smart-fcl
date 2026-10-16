"""Offline readers; preserve page/table boundaries. No document instructions executed."""
import csv, io, os, shutil, subprocess, re, statistics, gc, sys, ctypes, zipfile, time, xml.etree.ElementTree as ET
from pathlib import Path
try:
    import fitz
except ImportError:
    fitz = None
from PIL import Image, ImageOps, ImageStat
from docx import Document
from openpyxl import load_workbook
import xlrd, pytesseract
from .db import ROOT

# Tesseract's OpenMP defaults can oversubscribe a small Render instance and time out.
os.environ.setdefault('OMP_THREAD_LIMIT', '1')

if (ROOT / 'tools/tessdata/eng.traineddata').exists():
    os.environ['TESSDATA_PREFIX'] = str(ROOT / 'tools/tessdata')
elif (ROOT / 'tools/ocr/usr/share/tesseract-ocr/5/tessdata/eng.traineddata').exists():
    os.environ['TESSDATA_PREFIX'] = str(ROOT / 'tools/ocr/usr/share/tesseract-ocr/5/tessdata')

SUPPORTED = {'.pdf', '.png', '.jpg', '.jpeg', '.tif', '.tiff', '.xlsx', '.xls', '.csv', '.txt', '.docx'}
MAX_BYTES = 25 * 1024 * 1024
MAX_PAGES = 100
Image.MAX_IMAGE_PIXELS = 35_000_000

def ocr_status():
    cloud_ocr = ROOT / 'tools/ocr/tesseract'
    candidate = os.getenv('TESSERACT_CMD') or (str(cloud_ocr) if cloud_ocr.exists() else shutil.which('tesseract'))
    for p in [ROOT / 'tools/tesseract/tesseract.exe', Path('C:/Program Files/Tesseract-OCR/tesseract.exe')]:
        if not candidate and p.exists():
            candidate = str(p)
    if not candidate:
        return {'ready': False, 'languages': [], 'detail': 'Chưa cài Tesseract OCR'}
    try:
        r = subprocess.run(
            [candidate, '--list-langs'],
            capture_output=True,
            text=True,
            timeout=8,
            creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0),
        )
        langs = [s.strip() for s in r.stdout.splitlines()[1:] if s.strip()]
        return {'ready': r.returncode == 0 and 'eng' in langs, 'languages': langs, 'detail': 'Tesseract local', 'cmd': candidate}
    except Exception:
        return {'ready': False, 'languages': [], 'detail': 'Không chạy được Tesseract'}

def ocr_tables(data):
    """Recover simple aligned tables from word coordinates, never infer missing cells."""
    from .extraction import ITEM_ALIASES, CONT_ALIASES
    words = []
    for i, text in enumerate(data['text']):
        if text.strip() and not re.fullmatch(r'[|¦]+', text.strip()) and float(data['conf'][i]) >= 0:
            words.append({'text': text.strip(), 'x': data['left'][i], 'y': data['top'][i] + data['height'][i]/2,
                          'right': data['left'][i]+data['width'][i], 'height': data['height'][i]})
    if not words:
        return []
    tolerance = max(4, statistics.median(w['height'] for w in words) * .5)
    rows = []
    for word in sorted(words, key=lambda w: (w['y'], w['x'])):
        if rows and abs(word['y'] - rows[-1][0]['y']) < tolerance:
            rows[-1].append(word)
        else:
            rows.append([word])
    rows = [sorted(row, key=lambda w: w['x']) for row in rows]
    compact = lambda s: re.sub(r'[^a-z0-9]', '', s.casefold())
    tables = []
    for row_index, row in enumerate(rows):
        for mode, aliases, required in [('item', ITEM_ALIASES, {'description','quantity'}), ('container', CONT_ALIASES, {'container_number'})]:
            columns = []; cursor = 0
            while cursor < len(row):
                matches = []
                for length in range(1, min(4, len(row)-cursor)+1):
                    phrase = compact(' '.join(w['text'] for w in row[cursor:cursor+length]))
                    for key, labels in aliases.items():
                        if phrase in {compact(label) for label in labels}:
                            matches.append((length, key))
                if not matches:
                    cursor += 1
                    continue
                length, key = max(matches)
                columns.append((key, row[cursor]['x']))
                cursor += length
            if not required.issubset({col[0] for col in columns}) or len(columns) < 2:
                continue
            if len({col[0] for col in columns}) != len(columns):
                continue
            table = [[aliases[key][0] for key, _ in columns]]
            for data_row in rows[row_index+1:]:
                cells = [''] * len(columns)
                for word in data_row:
                    center = (word['x']+word['right'])/2
                    index = max((i for i, (_, x) in enumerate(columns) if center >= x-tolerance), default=0)
                    cells[index] = (cells[index]+' '+word['text']).strip()
                values = {key:cells[i] for i, (key, _) in enumerate(columns)}
                if mode == 'container':
                    valid = bool(re.fullmatch(r'[A-Z]{4}\d{7}', re.sub(r'\s','',values.get('container_number','')).upper()))
                else:
                    valid = bool(values.get('description') and re.fullmatch(r'[\d., ]+', values.get('quantity','')))
                if not valid:
                    break
                table.append(cells)
            if len(table) > 1:
                tables.append(table)
    return tables + ocr_packing_summary(data) + ocr_labelled_weights(data)


def ocr_labelled_weights(data):
    """Retain the spatial link when a form label/value occupy different OCR blocks."""
    words=[{'text':t.strip(),'x':data['left'][i],'right':data['left'][i]+data['width'][i],
            'y':data['top'][i]+data['height'][i]/2,'h':data['height'][i]}
           for i,t in enumerate(data['text']) if t.strip() and float(data['conf'][i])>=0]
    rows=[]
    for label in words:
        if not re.fullmatch(r'(?i)\(?gross\)?',label['text']):continue
        same=[w for w in words if w['x']>label['right'] and abs(w['y']-label['y'])<max(6,label['h']*.6)]
        units=[w for w in same if re.fullmatch(r'(?i)KGS?|KGM|MT|LBS?',w['text'])]
        numbers=[w for w in same if re.fullmatch(r'\d[\d.,]*',w['text'])]
        if len(units)==len(numbers)==1 and numbers[0]['x']<units[0]['x']:
            rows.append([label['text'],numbers[0]['text'],units[0]['text']])
    return [[['SOURCE_LABEL','SOURCE_VALUE','SOURCE_UNIT'],*rows]] if rows else []


def ocr_packing_summary(data):
    """Bind printed PL total cells by OCR coordinates, never calculate totals."""
    words=[{'text':t.strip(),'x':data['left'][i],'y':data['top'][i]+data['height'][i]/2,
            'h':data['height'][i]} for i,t in enumerate(data['text']) if t.strip() and float(data['conf'][i])>=0]
    if not words:return []
    tolerance=max(4,statistics.median(w['h'] for w in words)*.6)
    rows=[]
    for w in sorted(words,key=lambda w:(w['y'],w['x'])):
        if rows and abs(w['y']-rows[-1][0]['y'])<=tolerance:rows[-1].append(w)
        else:rows.append([w])
    for row in rows:row.sort(key=lambda w:w['x'])
    compact=lambda t:re.sub(r'[^a-z]','',t.casefold())
    for hi,row in enumerate(rows):
        names={compact(w['text']):w for w in row}
        if not {'net','gross'}.issubset(names):continue
        left=names.get('ctn',{}).get('x')
        net,gross=names['net']['x'],names['gross']['x']
        if not net<gross:continue
        units=[w for r in rows[hi+1:hi+7] for w in r if re.search(r'(?i)\bkgs?\b',w['text'])]
        if not any(net-20<=w['x']<gross for w in units) or not any(gross-20<=w['x'] for w in units):continue
        for ri in range(hi+2,len(rows)):
            footer=' '.join(w['text'] for w in rows[ri]).lstrip('[|')
            if not re.match(r'(?i)tare\s+weight\s*\(kgs?\)',footer):continue
            aggregate=rows[ri-1]
            if rows[ri][0]['y']-aggregate[0]['y']>tolerance*6:continue
            cells=[]
            for w in aggregate:
                raw=w['text'].strip('[]{}|')
                if not re.fullmatch(r'\d[\d.,]*',raw):break
                cells.append((w['x'],raw))
            else:
                nw=[v for x,v in cells if net-20<=x<gross-20]
                gw=[v for x,v in cells if gross-20<=x<gross+(gross-net)-20]
                # Quantity may be another column before Net; the carton count
                # must be the single value beside the CTN header's x-position.
                pkg=[v for x,v in cells if left is not None and abs(x-left)<80]
                if len(nw)==len(gw)==1:
                    if len(pkg)==1:
                        return [[['CTN','Description','Net Wt. (kgs)','Gross Wt. (kgs)'],
                                 [pkg[0],'',nw[0],gw[0]],[footer,'','','']]]
                    return [[['Description','Net Wt. (kgs)','Gross Wt. (kgs)'],
                             ['',nw[0],gw[0]],[footer,'','']]]
    return []

def ocr_grid_totals(image):
    """Read explicit total cells inside a ruled table; never sum or fill a cell.

    Whole-page OCR can merge bundles, pieces and cubic metres into one stream.
    Use printed grid boundaries and a literal Total label to retain their roles.
    Stamps/illegible cells remain empty, not calculated from the preceding rows.
    """
    status=ocr_status()
    if not status['ready']:return []
    pytesseract.pytesseract.tesseract_cmd=status['cmd']
    lang='eng+vie' if 'vie' in status['languages'] else 'eng'
    started=time.monotonic();calls=0
    def centres(values):
        groups=[]
        for value in values:
            if groups and value-groups[-1][-1]<4:groups[-1].append(value)
            else:groups.append([value])
        return [sum(group)//len(group) for group in groups]
    gray=ImageOps.grayscale(image)
    mask=gray.point(lambda v:255 if v<180 else 0)
    try:
        ys=centres([y for y,v in enumerate(mask.resize((1,gray.height),Image.Resampling.BOX).getdata()) if v>255*.65])
        rows=[]
        for top,bottom in zip(ys,ys[1:]):
            if not 20<=bottom-top<=gray.height*.22:continue
            with mask.crop((0,top+5,gray.width,bottom-5)) as band:
                xs=centres([x for x,v in enumerate(band.resize((gray.width,1),Image.Resampling.BOX).getdata()) if v>255*.75])
            if 5<=len(xs)<=15 and all(b-a>=30 for a,b in zip(xs,xs[1:])):rows.append((top,bottom,xs))
        def cell(row,index):
            nonlocal calls
            if calls>=40 or time.monotonic()-started>25:raise ValueError('Hết ngân sách OCR ô tổng')
            calls+=1
            top,bottom,xs=row
            with gray.crop((xs[index]+5,top+5,xs[index+1]-5,bottom-5)) as crop:
                return pytesseract.image_to_string(crop,lang=lang,config='--psm 6',timeout=5).strip()[:250]
        for hi,row in enumerate(rows[:18]):
            first=[cell(row,i) for i in range(min(2,len(row[2])-1))]
            if not re.search(r'(?i)description\s+of\s+goods|commodity', ' '.join(first)):continue
            header=first+[cell(row,i) for i in range(2,len(row[2])-1)]
            if not re.search(r'(?i)quantity|weight|no[. ]*of', ' '.join(header)):continue
            for footer in list(reversed(rows[hi+1:]))[:6]:
                if len(footer[2])!=len(row[2]) or any(abs(a-b)>12 for a,b in zip(row[2],footer[2])):continue
                label=cell(footer,0)
                if not re.fullmatch(r'(?i)\s*(?:grand\s+)?total\s*[:：.]?\s*',label):continue
                values=[label]+[cell(footer,i) for i in range(1,len(footer[2])-1)]
                # Only literal numeric footer cells are usable as totals.
                values=[values[0]]+[v if re.fullmatch(r'\d[\d., ]*',v) else '' for v in values[1:]]
                if any(values[1:]):return [[header,values]]
            break
    except (ValueError,RuntimeError):
        return []  # The original full-page OCR remains available.
    finally:
        gray.close();mask.close()
    return []


def prepare_ocr_image(image):
    image = ImageOps.exif_transpose(image).convert('RGB')
    # Small phone captures are easier to read after bounded enlargement. Never invent text.
    if image.width < 1600:
        scale=min(1600/image.width,(35_000_000/(image.width*image.height))**.5)
        if scale>1:
            image=image.resize((int(image.width*scale),int(image.height*scale)),Image.Resampling.LANCZOS)
    return image

def release_reader_memory():
    """Discard native PDF/image caches between pages on memory-limited workers."""
    if fitz is not None:
        fitz.TOOLS.store_shrink(100)
    gc.collect()
    if sys.platform.startswith('linux'):
        try:
            ctypes.CDLL(None).malloc_trim(0)
        except (AttributeError, OSError):
            pass  # Non-glibc platforms can still release Python/PDF objects.


def pdf_ocr_image(page):
    # OCR already reads grayscale. Avoid simultaneous RGB pixmap, PNG buffer,
    # decoded RGB image and grayscale copies of a full page.
    pix = page.get_pixmap(matrix=fitz.Matrix(200/72, 200/72), colorspace=fitz.csGRAY, alpha=False)
    return Image.frombytes('L', (pix.width, pix.height), pix.samples)


def ocr(image, include_tables=False, psm=3):
    try:
        return _ocr(image, include_tables, psm)
    finally:
        release_reader_memory()


def _ocr(image, include_tables=False, psm=3):
    status = ocr_status()
    if not status['ready']:
        raise ValueError('OCR local chưa sẵn sàng. Cài Tesseract và thử lại; PDF text vẫn đọc được.')
    pytesseract.pytesseract.tesseract_cmd = status['cmd']
    image = ImageOps.exif_transpose(image).convert('RGB')
    if image.width * image.height > 35_000_000:
        raise ValueError('Ảnh quá lớn để OCR an toàn')
    lang = 'eng+vie' if 'vie' in status['languages'] else 'eng'
    original_width=image.width
    image=prepare_ocr_image(image)
    gray = ImageOps.autocontrast(ImageOps.grayscale(image))
    # Automatic page segmentation preserves large headings and multiple text/table blocks.
    try:
        data = pytesseract.image_to_data(gray, lang=lang, config='--psm '+str(psm), output_type=pytesseract.Output.DICT, timeout=60)
    except RuntimeError as exc:
        if 'timeout' in str(exc).lower():
            raise ValueError('OCR vượt 60 giây cho một trang. Hãy dùng bản scan rõ hơn hoặc chia nhỏ chứng từ.') from None
        raise
    lines = {}
    scores = []
    for i, t in enumerate(data['text']):
        if not t.strip():
            continue
        key = (data['block_num'][i], data['par_num'][i], data['line_num'][i])
        lines.setdefault(key, []).append(t)
        if float(data['conf'][i]) >= 0:
            scores.append(float(data['conf'][i]))
    text = '\n'.join(' '.join(v) for v in lines.values())
    quality = []
    if not text.strip():
        raise ValueError('OCR không đọc được chữ. Hãy dùng bản scan rõ hơn.')
    if not scores or sum(scores) / len(scores) < 70:
        quality.append('OCR chất lượng thấp; kiểm tra từng giá trị với bản gốc.')
    if original_width < 800 or ImageStat.Stat(gray).stddev[0] < 20:
        quality.append('Ảnh nhỏ/độ tương phản thấp; có thể bỏ sót chữ.')
    if 'vie' not in status['languages']:
        quality.append('Chưa cài ngôn ngữ vie; dấu tiếng Việt có thể đọc sai.')
    unclear=[data['text'][i].strip() for i in range(len(data['text']))
             if 0<=float(data['conf'][i])<70 and re.fullmatch(r'[A-Za-z0-9/-]*\d[A-Za-z0-9/-]{2,}',data['text'][i].strip())]
    if unclear:quality.append('Ký tự mã/số chưa rõ trong lượt này: '+', '.join(dict.fromkeys(unclear[:20]))+'.')
    result = (text, ' '.join(quality) or 'OCR cần nhân viên kiểm tra')
    return (*result, ocr_tables(data)) if include_tables else result

def decode(data):
    for encoding in ['utf-8-sig', 'utf-16', 'cp1258']:
        try:
            return data.decode(encoding)
        except UnicodeError:
            pass
    raise ValueError('Không đọc được encoding. Lưu file dưới UTF-8 và thử lại.')


def read_pdf_image(image):
    """Retain genuinely blank pages without hiding unreadable content pages."""
    gray = ImageOps.grayscale(image)
    hist = gray.histogram()
    pixels = gray.width * gray.height
    if sum(hist[:220]) / pixels < .00001 and ImageStat.Stat(gray).stddev[0] < 1:
        return '', 'Blank page', 'Trang trắng được giữ nguyên; kiểm tra bản gốc.', []
    text, quality, tables = ocr(image, include_tables=True)
    # Forms with disconnected cells can lose values under automatic block segmentation.
    # Add a sparse pass only for form/table pages without recovered tables; keep both verbatim.
    if not tables and re.search(r'(?i)\b(?:quantity|qty|container|bill\s+of\s+lading|packing\s+list|số\s+tờ\s+khai|người\s+xuất\s+khẩu)\b',text):
        try:
            sparse, sparse_quality, sparse_tables = ocr(image, include_tables=True, psm=11)
            if sparse.strip() != text.strip():
                text += '\n\n[OCR sparse layer — giữ nguyên hai lượt đọc, đối chiếu ảnh nếu khác nhau]\n'+sparse
                quality += ' Có lượt OCR bổ sung cho các ô rời; đối chiếu ảnh gốc nếu hai lượt khác nhau.'
                quality += ' Lượt ô rời: '+sparse_quality
                tables = tables or sparse_tables
        except ValueError as exc:
            quality += ' OCR bổ sung chưa hoàn tất: '+str(exc)
    if re.search(r'(?i)vessel|voyage|chuyến|tên\s+tàu',text):
        try:
            line_text,line_quality,line_tables=ocr(image,include_tables=True,psm=6)
            if line_text.strip()!=text.strip():
                text+='\n\n[OCR line layer — đọc lại trực tiếp các dòng trên cùng ảnh]\n'+line_text
                quality+=' Lượt dòng: '+line_quality
        except ValueError as exc:
            quality+=' OCR dòng bổ sung chưa hoàn tất: '+str(exc)
    if re.search(r'(?i)packing\s+list|invoice|sales?\s+contract',text):
        grid=ocr_grid_totals(image)
        if grid:
            tables+=grid
            text+='\n\n[OCR grid cells — nhãn cột và ô Total đọc trực tiếp từ cùng ảnh]\n'+'\n'.join(' | '.join(row) for table in grid for row in table)
            quality+=' Đã giữ liên kết nhãn cột và ô tổng theo đường kẻ bảng; ô không đọc rõ để trống.'
    quality+=' Đã đọc bổ sung ô/dòng theo bố cục.'
    return text, 'OCR', quality, tables


def hybrid_text(text, image, tables):
    """Keep selectable text AND image-only labels. Neither layer overwrites the source."""
    status = ocr_status()
    if not status['ready']:
        return text, 'PDF text', 'PDF có ảnh; OCR chưa sẵn sàng, có thể thiếu nhãn/tiêu đề.', tables
    try:
        raster_text, quality, raster_tables = ocr(image, include_tables=True)
    except ValueError as exc:
        return text, 'PDF text', 'Không đọc thêm được lớp ảnh: '+str(exc), tables
    return (text+'\n\n[OCR image layer — đối chiếu ảnh gốc nếu khác lớp chữ]\n'+raster_text,
            'PDF text + OCR', quality, tables or raster_tables)

def render_pdf_page_png(path, page_number, scale=1.5):
    path = Path(path)
    if fitz is not None:
        with fitz.open(path) as pdf:
            if page_number > len(pdf):
                raise ValueError('Không có trang này')
            pix = pdf[page_number - 1].get_pixmap(matrix=fitz.Matrix(scale, scale))
            return pix.tobytes('png')
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(path))
    try:
        if page_number > len(pdf):
            raise ValueError('Không có trang này')
        pil_img = pdf[page_number - 1].render(scale=scale).to_pil()
        buf = io.BytesIO()
        pil_img.save(buf, format='PNG')
        return buf.getvalue()
    finally:
        pdf.close()

def preserve_pdf_content_order(page, text, tables):
    """Keep a second verbatim reading when geometric sorting interleaves table cells."""
    if not tables:
        return text
    original = page.get_text('text', sort=False)
    # No guessed column reconstruction, numeric cleanup or source overwriting.
    # Large documents keep the existing reading rather than silently truncating it.
    if (original.strip() and len(original)<=20000
            and ' '.join(original.split())!=' '.join(text.split())):
        return text+'\n\n[PDF content-order layer — same page, original text sequence]\n'+original
    return text


def read_document(path, progress=lambda *_: None, cancelled=lambda: False):
    path = Path(path)
    ext = path.suffix.lower()
    pages = []

    def add(text, method, quality='', tables=None):
        if cancelled():
            raise InterruptedError('Đã hủy xử lý')
        if len(text) > 250000:
            raise ValueError('Một trang quá dài; giới hạn 250.000 ký tự')
        pages.append({'number': len(pages) + 1, 'text': text, 'method': method, 'quality': quality, 'tables': tables or []})
        progress(len(pages))

    if ext == '.pdf':
        if fitz is not None:
            with fitz.open(path) as pdf:
                if pdf.is_encrypted:
                    raise ValueError('PDF có mật khẩu. Hãy mở khóa trước khi tải lên.')
                if len(pdf) > MAX_PAGES:
                    raise ValueError('PDF vượt giới hạn 100 trang')
                for page in pdf:
                    if cancelled():
                        raise InterruptedError('Đã hủy xử lý')
                    text = page.get_text('text', sort=True)
                    tables = []
                    if len(text.strip()) >= 20:
                        try:
                            tables = [t.extract() for t in page.find_tables().tables]
                        except Exception:
                            pass
                        text = preserve_pdf_content_order(page, text, tables)
                        if page.get_images():
                            with pdf_ocr_image(page) as image:
                                add(*hybrid_text(text, image, tables))
                        else:
                            add(text, 'PDF text', '', tables)
                    else:
                        with pdf_ocr_image(page) as image:
                            text, method, quality, tables = read_pdf_image(image)
                        add(text, method, quality, tables)
                    release_reader_memory()
        else:
            from pypdf import PdfReader
            import pdfplumber
            import pypdfium2 as pdfium
            try:
                reader_pdf = PdfReader(str(path))
            except Exception as exc:
                raise ValueError('Không đọc được tệp PDF: ' + str(exc))
            if reader_pdf.is_encrypted:
                raise ValueError('PDF có mật khẩu. Hãy mở khóa trước khi tải lên.')
            if len(reader_pdf.pages) > MAX_PAGES:
                raise ValueError('PDF vượt giới hạn 100 trang')
            with pdfplumber.open(str(path)) as plumber_pdf:
                pdfium_doc = pdfium.PdfDocument(str(path))
                try:
                    for idx, p_page in enumerate(plumber_pdf.pages):
                        if cancelled():
                            raise InterruptedError('Đã hủy xử lý')
                        text = p_page.extract_text() or ''
                        tables = []
                        if len(text.strip()) >= 20:
                            try:
                                tables = p_page.extract_tables() or []
                            except Exception:
                                pass
                            if p_page.images:
                                add(*hybrid_text(text, pdfium_doc[idx].render(scale=200/72).to_pil(), tables))
                            else:
                                add(text, 'PDF text', '', tables)
                        else:
                            pil_img = pdfium_doc[idx].render(scale=200/72).to_pil()
                            text, method, quality, tables = read_pdf_image(pil_img)
                            add(text, method, quality, tables)
                finally:
                    pdfium_doc.close()
    elif ext in ['.png', '.jpg', '.jpeg', '.tif', '.tiff']:
        with Image.open(path) as im:
            count = getattr(im, 'n_frames', 1)
            if count > MAX_PAGES:
                raise ValueError('Ảnh vượt giới hạn 100 trang')
            for n in range(count):
                im.seek(n)
                text, q, tables = ocr(im.copy(), include_tables=True)
                add(text, 'OCR', q, tables)
    elif ext == '.xlsx':
        # VNACCS exports format empty cells up to column IV. Count only cells
        # containing values/formulas, without dropping genuine distant data.
        bounds={}
        from openpyxl.utils.cell import column_index_from_string
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist())>80*1024*1024:raise ValueError('Workbook vượt giới hạn dữ liệu giải nén')
            for name in archive.namelist():
                if not re.fullmatch(r'xl/worksheets/sheet\d+\.xml',name):continue
                last_row,last_col=1,1
                with archive.open(name) as source:
                    for _,cell in ET.iterparse(source,events=['end']):
                        if cell.tag.rsplit('}',1)[-1]!='c':continue
                        populated=any(child.tag.rsplit('}',1)[-1] in ['v','is','f'] and (child.text or list(child)) for child in cell)
                        if populated:
                            ref=re.fullmatch(r'([A-Z]+)(\d+)',cell.get('r',''))
                            if ref:last_col=max(last_col,column_index_from_string(ref[1]));last_row=max(last_row,int(ref[2]))
                        cell.clear()
                if last_row>10000 or last_col>100:raise ValueError('Sheet vượt 10.000 dòng hoặc 100 cột có dữ liệu')
                bounds[name]=(last_row,last_col)
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            if len(wb.worksheets) > MAX_PAGES:
                raise ValueError('Workbook có quá nhiều sheet')
            for sheet in wb:
                max_row,max_col=bounds.get(sheet._worksheet_path,(sheet.max_row,sheet.max_column))
                if max_row>10000 or max_col>100:raise ValueError('Sheet vượt giới hạn dữ liệu')
                table = [[str(v) if v is not None else '' for v in row] for row in sheet.iter_rows(max_row=max_row,max_col=max_col,values_only=True)]
                add(
                    sheet.title + '\n' + '\n'.join(' | '.join(r) for r in table),
                    'Excel',
                    'Đơn vị hiển thị là Sheet bảng tính (không phải trang in cố định). Công thức chỉ được đọc dưới dạng text, không chạy.',
                    [table],
                )
        finally:
            wb.close()
    elif ext == '.xls':
        wb = xlrd.open_workbook(path)
        for s in wb.sheets():
            if s.nrows > 10000 or s.ncols > 100:
                raise ValueError('XLS vượt giới hạn bảng')
            table = [[str(v) for v in s.row_values(i)] for i in range(s.nrows)]
            add(
                s.name + '\n' + '\n'.join(' | '.join(r) for r in table),
                'Excel',
                'Đơn vị hiển thị là Sheet bảng tính (không phải trang in cố định).',
                [table],
            )
    elif ext == '.docx':
        d = Document(path)
        tables = [[[cell.text for cell in row.cells] for row in t.rows] for t in d.tables]
        text = '\n'.join(p.text for p in d.paragraphs) + '\n' + '\n'.join(' | '.join(row) for table in tables for row in table)
        add(text, 'Word', 'Đơn vị hiển thị Word là khối nội dung logic; không phải số trang in cố định.', tables)
    elif ext in ['.txt', '.csv']:
        text = decode(path.read_bytes())
        tables = []
        if ext == '.csv':
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
                tables = [list(csv.reader(io.StringIO(text), dialect))]
            except csv.Error:
                pass
        add(text, 'Text', '', tables)
    else:
        raise ValueError('Định dạng không hỗ trợ')
    if not pages or not any(p['text'].strip() for p in pages):
        raise ValueError('Tệp không có nội dung đọc được')
    return pages
