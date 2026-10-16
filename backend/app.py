import hashlib, io, json, threading, time, zipfile, sqlite3, os
from pathlib import Path
from typing import Literal, Optional, List, Union, Dict, Any
from urllib.parse import urlparse
from fastapi import FastAPI, UploadFile, File, HTTPException, Request
from fastapi.responses import FileResponse, Response, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from fastapi.exceptions import RequestValidationError
from pydantic import BaseModel, Field
try:
    from pydantic import ConfigDict
except ImportError:
    ConfigDict = None
from PIL import Image, ImageOps
from openpyxl import load_workbook
from . import db, reader, extraction, validation, ai, excel, templates, pdf_bundle, source_locator, consistency
from .dictionary import COMPARISON_TYPES
from . import pipeline
from .dictionary import FIELDS, TYPES, default_requirements, commercial_profile, REVIEW_GROUPS, REVIEW_FIELDS, PROFILE_TYPES, PROFILE_VARIANTS, DOCUMENT_LABELS
from .normalize import normalize

db.migrate()
app = FastAPI(title='U&I SMART', version='1.1.0')
app.add_middleware(TrustedHostMiddleware, allowed_hosts=['127.0.0.1', 'localhost', 'testserver'])
JOBS = {}
LOCK = threading.RLock()

@app.middleware('http')
async def local_guard(request: Request, call_next):
    if request.method not in ['GET', 'HEAD', 'OPTIONS']:
        origin = request.headers.get('origin')
        if origin and (urlparse(origin).netloc != request.headers.get('host') or urlparse(origin).scheme not in ['http', 'https']):
            return JSONResponse({'detail': 'Origin không được phép'}, 403)
        if request.headers.get('x-ui-smart') != 'local':
            return JSONResponse({'detail': 'Thiếu header bảo vệ thao tác local'}, 403)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['Content-Security-Policy'] = (
        "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; "
        "img-src 'self' data: blob:; font-src 'self'; connect-src 'self'; "
        "frame-ancestors 'none'; object-src 'none'"
    )
    response.headers['Cache-Control'] = 'no-store' if request.url.path.startswith('/api/') else 'no-cache'
    return response

@app.exception_handler(ValueError)
async def invalid(_request, exc):
    return JSONResponse({'detail': str(exc)[:600]}, 422)

@app.exception_handler(RequestValidationError)
async def invalid_request(_request, exc):
    # Do not echo raw validation inputs, especially API secrets.
    return JSONResponse({'detail': [{'type': e['type'], 'loc': list(e['loc']), 'msg': e['msg']} for e in exc.errors()]}, 422)

class Input(BaseModel):
    if ConfigDict is not None:
        model_config = ConfigDict(extra='forbid')
    else:
        class Config:
            extra = 'forbid'

    if not hasattr(BaseModel, 'model_dump'):
        def model_dump(self):
            return self.dict()

class CaseInput(Input):
    name: str = Field(min_length=1, max_length=160)
    customer: str = Field(default='', max_length=160)
    mode: Literal['FCL', 'AIR'] = 'FCL'

class ProcessInput(Input):
    mode: Literal['local', 'openai'] = 'local'
    allow_external: bool = False

class ClassInput(Input):
    type: Literal[
        'Invoice', 'Invoice & Packing List', 'Packing List', 'Booking', 'B/L', 'Sea Waybill', 'AWB', 'C/O',
        'Arrival Notice', 'Delivery Order', 'Contract', 'Purchase Order',
        'Insurance', 'Freight Invoice', 'Import Permit',
        'Inspection Certificate', 'Phytosanitary Certificate', 'Customs Declaration','Customs Release','Warehouse Receipt','Other'
    ]

class CollectionInput(Input):
    collection: Literal['document', 'reference']
    reason: str = Field(default='', max_length=1000)
    document_type: Optional[str] = None

class SplitGroup(ClassInput):
    start: int = Field(ge=1, le=100)
    end: int = Field(ge=1, le=100)

class SplitInput(Input):
    groups: List[SplitGroup] = Field(min_length=2, max_length=30)

class EditInput(Input):
    value: Union[str, Dict[str, Any]]
    note: str = Field(default='', max_length=1000)
    version: int = Field(ge=1)

class VerifyInput(Input):
    status: Literal['verified', 'rejected', 'pending']
    note: str = Field(default='', max_length=1000)
    version: int = Field(ge=1)

class NewRecord(Input):
    kind: Literal['field', 'item', 'container']
    key: str
    value: Union[str, Dict[str, Any]]
    note: str = Field(min_length=3, max_length=1000)
    document_id: Optional[str] = None
    page: Optional[int] = Field(default=None, ge=1, le=100)
    confirmed: bool = False

class ResolveInput(Input):
    note: str = Field(min_length=5, max_length=1000)

class CheckReviewInput(Input):
    note: str = Field(default='',max_length=1000)
    confirmed: bool = False
    versions: Dict[str,int]

class DiagnosticReviewInput(Input):
    note: str = Field(default='',max_length=1000)
    confirmed: bool = False
    token: str
    action: Literal['accept','exclude']
    value: Optional[Union[str,Dict[str,Any]]] = None

class RequirementInput(Input):
    required: bool
    note: str = Field(min_length=3, max_length=1000)

class ExportInput(Input):
    mode: Literal['Draft', 'Reviewed'] = 'Draft'
    workbook_type: Literal['A', 'B'] = 'A'
    selected: Optional[List[str]] = None

class AIInput(Input):
    api_key: str = Field(default='', max_length=500)
    model: str = Field(default='gpt-4.1-mini', max_length=100)

class TemplateInput(Input):
    case_id: str
    mapping: List[Dict[str, Any]]

def require_case(id):
    case = db.one('SELECT * FROM cases WHERE id=?', (id,))
    if not case:
        raise HTTPException(404, 'Không tìm thấy hồ sơ')
    return case

def require_doc(id):
    doc = db.one('SELECT * FROM documents WHERE id=?', (id,))
    if not doc:
        raise HTTPException(404, 'Không tìm thấy tài liệu')
    return doc

def require_record(id):
    record = db.one('SELECT * FROM records WHERE id=?', (id,))
    if not record:
        raise HTTPException(404, 'Không tìm thấy dữ liệu')
    if record['document_id']:
        require_working_doc(record['document_id'])
    return validation.unpack(record)

def require_working_doc(id):
    doc = require_doc(id)
    if doc.get('collection', 'document') != 'document':
        raise HTTPException(409, 'Tệp thuộc Giấy tờ khác. Chọn đúng loại và đưa vào bộ chính trước khi trích xuất hoặc xác minh.')
    return doc

def working_records(id):
    records = [validation.unpack(r) for r in db.rows("SELECT * FROM records WHERE case_id=? AND (document_id IS NULL OR document_id IN (SELECT id FROM documents WHERE collection='document')) ORDER BY rowid", (id,))]
    case = require_case(id)
    return [r for r in records if r['kind'] != 'container' and (r['kind'] != 'field' or r['key'] in REVIEW_FIELDS)] if commercial_profile(json.loads(case['requirements']),case['mode']) else records

def busy(case_id):
    if any(j['case_id'] == case_id and j['status'] == 'running' for j in JOBS.values()):
        raise HTTPException(409, 'Hồ sơ đang xử lý. Hủy/chờ hoàn tất trước khi chỉnh sửa.')

def bundle(id):
    case = require_case(id)
    docs = db.rows("SELECT * FROM documents WHERE case_id=? AND collection='document' ORDER BY rowid", (id,))
    for d in docs:
        d.pop('path')
        diagnostics=json.loads(d.get('reading_diagnostics') or '[]')
        if diagnostics and all(e.get('review') for e in diagnostics):
            d['reviewed_reading_warning']=d['error'];d['error']=''
        d['pages'] = db.rows('SELECT number,method,quality,length(text) characters FROM pages WHERE document_id=? ORDER BY number', (d['id'],))
        if d['ext'] == '.pdf' and len(d['pages']) > 1:
            text_pages = db.rows('SELECT number,text FROM pages WHERE document_id=? ORDER BY number', (d['id'],))
            d['split_suggestion'] = pdf_bundle.suggest(text_pages)
            d['mixed_pdf'] = pdf_bundle.is_mixed(text_pages)
        d['source'] = db.one('SELECT s.id,s.name,s.hash,m.start_page,m.end_page FROM document_sources m JOIN bundle_sources s ON s.id=m.source_id WHERE m.document_id=?', (d['id'],))
    references = db.rows("SELECT id,name,type,ext,size,collection_reason,classification_evidence,reading_method,status,error FROM documents WHERE case_id=? AND collection='reference' ORDER BY rowid", (id,))
    for d in references:
        d['pages'] = db.rows('SELECT number,method,quality,length(text) characters FROM pages WHERE document_id=? ORDER BY number', (d['id'],))
    records = working_records(id)
    checks = db.rows('SELECT * FROM checks WHERE case_id=? ORDER BY rowid', (id,))
    for c in checks:
        c['detail'] = json.loads(c.pop('detail_json'))
        c['title'] = validation.human_check_label(c['key'], {d['id']: d for d in docs})
    checks=[c for c in checks if not c['key'].startswith('ai-semantic:') or consistency.semantic_difference(c,records)]
    case['requirements'] = json.loads(case['requirements'])
    return {
        'case': case,
        'documents': docs,
        'reference_files': references,
        'other_files': references,
        'records': records,
        'checks': checks,
        'consistency':consistency.report(case,docs,records,checks),
        'blockers': validation.blockers(case, docs, records, checks),
        'exports': db.rows('SELECT id,created_at,version,mode,workbook_type,case_revision FROM exports WHERE case_id=? ORDER BY version DESC', (id,)),
        'active_job': next((j for j in JOBS.values() if j['case_id'] == id and j['status'] == 'running'), None),
    }

def revalidate(c, id):
    # Remove an obsolete parser warning from saved pending dates, including a
    # staff-entered date. Keep raw source, method, note and approval state.
    for row in c.execute("SELECT * FROM records WHERE case_id=? AND kind='field' AND status='pending'",(id,)).fetchall():
        record=validation.unpack(row)
        if FIELDS.get(record['key'],{}).get('type')!='date' or not str(record['value'].get('warning','')).startswith('Định dạng ngày chưa rõ'):continue
        value=normalize(record['value'].get('text',''),'date')
        if value['warning']:continue
        c.execute('UPDATE records SET value_json=?,version=version+1 WHERE id=?',(db.encode({**record['value'],**value}),record['id']))
        db.audit(c,id,'Cập nhật định dạng ngày đã đọc rõ',record['id'],record['value'],value,actor='Hệ thống')
    case = dict(c.execute('SELECT * FROM cases WHERE id=?', (id,)).fetchone())
    docs = [dict(r) for r in c.execute("SELECT * FROM documents WHERE case_id=? AND collection='document'", (id,))]
    records = [validation.unpack(r) for r in c.execute("SELECT * FROM records WHERE case_id=? AND (document_id IS NULL OR document_id IN (SELECT id FROM documents WHERE collection='document'))", (id,))]
    checks=validation.persist_checks(c, case, docs, records)
    if commercial_profile(json.loads(case['requirements']),case['mode']):
        from .acceptance import apply
        apply(c,id)
    return checks

def add_record(c, case_id, document_id, r):
    case = dict(c.execute('SELECT * FROM cases WHERE id=?',(case_id,)).fetchone())
    if (r['kind'] == 'container' or (r['kind'] == 'field' and r['key'] not in REVIEW_FIELDS)) and commercial_profile(json.loads(case['requirements']),case['mode']):
        return
    c.execute(
        'INSERT INTO records VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
        (db.uid(), case_id, document_id, r['kind'], r['key'], db.encode(r['raw']), db.encode(r['value']), r['page'], r['evidence'], r['method'], 'pending', '', db.now(), 1),
    )

def replace_extracted_records(c, case_id, document_id, records):
    """Reconcile machine records; identical source/value keeps its ID and review links."""
    existing = [validation.unpack(r) for r in c.execute('SELECT * FROM records WHERE document_id=? ORDER BY rowid', (document_id,))]
    protected = [r for r in existing if r['status'] not in ('pending', 'auto_accepted') or r['method'].startswith('Manual')]
    automated = [r for r in existing if r not in protected]
    def identity(r):
        return pipeline.digest({k:r[k] for k in ['kind','key','page','raw','value','evidence','method']})
    kept = set()
    for r in records:
        if any(p['kind']==r['kind'] and p['key']==r['key'] and p['page']==r['page'] and p['raw']==r['raw'] for p in protected):
            continue
        match = next((old for old in automated if old['id'] not in kept and identity(old)==identity(r)), None)
        if match:
            kept.add(match['id'])
        else:
            add_record(c, case_id, document_id, r)
    for old in automated:
        if old['id'] not in kept:
            c.execute('DELETE FROM records WHERE id=?', (old['id'],))

@app.get('/api/health')
def health():
    try:
        db.one('SELECT version FROM schema_versions')
        ready = True
    except Exception:
        ready = False
    ocr = reader.ocr_status()
    ocr.pop('cmd', None)
    return {
        'ok': ready,
        'deployment': {'cloud': os.getenv('UI_SMART_CLOUD') == '1', 'ephemeral': os.getenv('UI_SMART_CLOUD') == '1'},
        'database_ready': ready,
        'service': 'U&I SMART',
        'version': '1.1.0',
        'ocr': ocr,
        'ai': ai.status(),
        'supported': sorted(reader.SUPPORTED),
        'max_bytes': reader.MAX_BYTES,
    }

@app.get('/api/dictionary')
def dictionary():
    return {'fields': FIELDS, 'document_types': COMPARISON_TYPES, 'document_labels':DOCUMENT_LABELS, 'review_groups':REVIEW_GROUPS, 'review_fields':REVIEW_FIELDS, 'legacy_document_types':TYPES, 'extraction_rules_version':10}

@app.get('/api/cases')
def list_cases():
    return db.rows(
        "SELECT c.*,(SELECT count(*) FROM documents d WHERE d.case_id=c.id AND d.collection='document') document_count,"
        "(SELECT count(*) FROM records r WHERE r.case_id=c.id AND r.status='pending' AND (r.document_id IS NULL OR r.document_id IN (SELECT id FROM documents WHERE collection='document'))) pending_count,"
        "(SELECT count(*) FROM exports e WHERE e.case_id=c.id) export_count "
        "FROM cases c ORDER BY updated_at DESC"
    )

@app.post('/api/cases', status_code=201)
def new_case(body: CaseInput):
    id = 'FCL-' + db.uid()[:12] if body.mode == 'FCL' else 'AIR-' + db.uid()[:12]
    stamp = db.now()
    with db.connect() as c:
        c.execute(
            'INSERT INTO cases VALUES (?,?,?,?,?,?,?,?,?,?)',
            (id, body.name.strip(), body.customer.strip(), body.mode, 'Draft', stamp, stamp, None, 0, db.encode(default_requirements(body.mode))),
        )
        db.audit(c, id, 'Tạo hồ sơ', after=body.model_dump())
    return bundle(id)

@app.get('/api/cases/{id}')
def get_case(id: str):
    return bundle(id)

@app.delete('/api/cases/{id}')
def delete_case(id: str):
    with LOCK:
        busy(id)
        require_case(id)
        paths = db.rows('SELECT path FROM documents WHERE case_id=? UNION ALL SELECT path FROM exports WHERE case_id=? UNION ALL SELECT path FROM bundle_sources WHERE case_id=?', (id, id, id))
        with db.connect() as c:
            c.execute('DELETE FROM cases WHERE id=?', (id,))
            db.audit(c, id, 'Xóa hồ sơ', reason='Người dùng xác nhận xóa')
        for r in paths:
            p = Path(r['path']).resolve()
            if p.is_relative_to(db.DATA):
                p.unlink(missing_ok=True)
    return {'ok': True}

def validate_zip(path):
    if path.suffix in ['.docx', '.xlsx']:
        try:
            with zipfile.ZipFile(path) as z:
                if sum(i.file_size for i in z.infolist()) > 150 * 1024 * 1024:
                    raise ValueError('File Office giải nén vượt giới hạn 150 MB')
                if len(z.infolist()) > 10000:
                    raise ValueError('File Office có quá nhiều mục')
        except zipfile.BadZipFile:
            raise ValueError('File Office không hợp lệ')

@app.post('/api/cases/{id}/documents')
async def upload(id: str, files: List[UploadFile] = File(...)):
    return await upload_files(id, files, 'document')

@app.post('/api/cases/{id}/reference-files')
async def upload_references(id: str, files: List[UploadFile] = File(...)):
    return await upload_files(id, files, 'reference')

async def upload_files(id, files, collection):
    require_case(id)
    if len(files) > 30:
        raise HTTPException(413, 'Tối đa 30 file mỗi lần')
    results = []
    for f in files:
        path = None
        try:
            name = Path((f.filename or '').replace('\\', '/')).name[:180]
            ext = Path(name).suffix.lower()
            if ext not in reader.SUPPORTED:
                raise ValueError('Định dạng chưa hỗ trợ: ' + ext)
            if collection == 'reference' and ext not in {'.pdf', '.png', '.jpg', '.jpeg', '.tif', '.tiff'}:
                raise ValueError('Bộ ảnh nhận PDF hoặc ảnh JPG, PNG, TIFF. Chứng từ Word/Excel tải ở mục Chứng từ.')
            data = await f.read(reader.MAX_BYTES + 1)
            if len(data) > reader.MAX_BYTES:
                raise ValueError('Tệp vượt giới hạn 25 MB')
            if not data:
                raise ValueError('Tệp rỗng')
            digest = hashlib.sha256(data).hexdigest()
            with LOCK:
                busy(id)
                if db.one('SELECT id FROM documents WHERE case_id=? AND hash=? UNION ALL SELECT id FROM bundle_sources WHERE case_id=? AND hash=?', (id, digest, id, digest)):
                    results.append({'name': name, 'status': 'duplicate', 'error': 'Nội dung đã có trong hồ sơ. Dùng Chuyển bộ nếu tệp đang ở mục khác.'})
                    continue
                doc_id = db.uid()
                path = db.DATA / 'files' / (doc_id + ext)
                path.write_bytes(data)
                validate_zip(path)
                with db.connect() as c:
                    c.execute('INSERT INTO documents(id,case_id,name,ext,size,hash,path,collection) VALUES (?,?,?,?,?,?,?,?)', (doc_id, id, name, ext, len(data), digest, str(path), collection))
                    if collection == 'document':
                        db.touch(c, id)
                        revalidate(c, id)
                    db.audit(c, id, 'Tải chứng từ' if collection == 'document' else 'Tải ảnh vào bộ riêng', doc_id, after={'name': name, 'size': len(data), 'hash': digest, 'collection': collection})
                results.append({'id': doc_id, 'name': name, 'status': 'uploaded'})
        except (ValueError, HTTPException) as exc:
            if path:
                path.unlink(missing_ok=True)
            results.append({'name': f.filename, 'status': 'error', 'error': str(exc.detail) if isinstance(exc, HTTPException) else str(exc)})
        finally:
            await f.close()
    return {'results': results, 'bundle': bundle(id)}

def move_collection(c, doc, target, reason):
    c.execute('UPDATE documents SET collection=?,collection_reason=? WHERE id=?', (target, reason, doc['id']))
    if target == 'document':
        c.execute("UPDATE documents SET type_verified=0 WHERE id=?", (doc['id'],))
        c.execute("UPDATE records SET status='pending',version=version+1 WHERE document_id=? AND status!='rejected'", (doc['id'],))
    db.touch(c, doc['case_id'])
    db.audit(c, doc['case_id'], 'Chuyển bộ tệp', doc['id'], before=doc.get('collection', 'document'), after=target, reason=reason)
    revalidate(c, doc['case_id'])

@app.patch('/api/documents/{id}/collection')
def collection(id: str, body: CollectionInput):
    with LOCK:
        doc = require_doc(id)
        busy(doc['case_id'])
        if body.document_type is not None and (body.collection != 'document' or body.document_type not in COMPARISON_TYPES):
            raise HTTPException(422, 'Chọn một trong 6 loại chứng từ chính trước khi đưa vào checklist.')
        if doc.get('collection', 'document') != body.collection:
            with db.connect() as c:
                move_collection(c, doc, body.collection, body.reason or 'Người dùng xác nhận chuyển bộ')
                if body.document_type is not None:
                    if doc['type'] != body.document_type:
                        c.execute("DELETE FROM records WHERE document_id=? AND status='pending' AND method NOT LIKE 'Manual%'", (id,))
                    c.execute("UPDATE documents SET type=?,type_verified=1,error=NULL,classification_evidence='',extraction_required=1 WHERE id=?", (body.document_type,id))
                    db.audit(c,doc['case_id'],'Xác nhận loại khi đưa vào bộ chính',id,before=doc['type'],after=body.document_type,reason=body.reason)
                    revalidate(c,doc['case_id'])
    return bundle(doc['case_id'])

@app.get('/api/documents/{id}/original')
def original_document(id: str):
    doc = require_doc(id)
    return FileResponse(doc['path'], filename=doc['name'])

@app.delete('/api/documents/{id}')
def delete_doc(id: str):
    with LOCK:
        doc = require_doc(id)
        busy(doc['case_id'])
        with db.connect() as c:
            c.execute('DELETE FROM documents WHERE id=?', (id,))
            if doc.get('collection', 'document') == 'document':
                db.touch(c, doc['case_id'])
            db.audit(c, doc['case_id'], 'Xóa tài liệu', id, before={'name': doc['name']})
            revalidate(c, doc['case_id'])
        Path(doc['path']).unlink(missing_ok=True)
    return bundle(doc['case_id'])

@app.patch('/api/documents/{id}/classification')
def classification(id: str, body: ClassInput):
    with LOCK:
        doc = require_working_doc(id)
        busy(doc['case_id'])
        pages = db.rows('SELECT number,text FROM pages WHERE document_id=? ORDER BY number', (id,))
        if doc['ext'] == '.pdf' and pdf_bundle.is_mixed(pages):
            raise HTTPException(409, 'PDF chứa nhiều loại chứng từ. Tách theo trang trước khi xác nhận loại.')
        with db.connect() as c:
            if doc['type'] != body.type:
                # Classification change requires extraction to be rerun, retain edited/verified records for review.
                c.execute("DELETE FROM records WHERE document_id=? AND status='pending' AND method NOT LIKE 'Manual%'", (id,))
                c.execute("UPDATE records SET status='pending',version=version+1,updated_at=? WHERE document_id=? AND status!='rejected'", (db.now(), id))
                c.execute('UPDATE documents SET extraction_required=1,classification_evidence=\'\' WHERE id=?',(id,))
            c.execute('UPDATE documents SET type=?,type_verified=1 WHERE id=?', (body.type, id))
            db.touch(c, doc['case_id'])
            db.audit(c, doc['case_id'], 'Xác nhận loại tài liệu', id, before=doc['type'], after=body.type)
            revalidate(c, doc['case_id'])
    return bundle(doc['case_id'])

@app.get('/api/sources/{id}/original')
def original_bundle(id: str):
    source = db.one('SELECT * FROM bundle_sources WHERE id=?', (id,))
    if not source:
        raise HTTPException(404, 'Không tìm thấy PDF gốc')
    return FileResponse(source['path'], filename=source['name'], media_type='application/pdf')

@app.post('/api/documents/{id}/split')
def split_pdf(id: str, body: SplitInput):
    from pypdf import PdfReader, PdfWriter
    created_paths = []
    with LOCK:
        doc = require_working_doc(id)
        busy(doc['case_id'])
        if doc['ext'] != '.pdf':
            raise HTTPException(422, 'Chỉ tách tệp PDF.')
        if db.one('SELECT document_id FROM document_sources WHERE document_id=?', (id,)):
            raise HTTPException(409, 'Chứng từ đã tách; dùng bản gốc trong hồ sơ mới để chia lại.')
        if db.one("SELECT id FROM records WHERE document_id=? AND (status!='pending' OR method LIKE 'Manual%')", (id,)):
            raise HTTPException(409, 'Chứng từ có dữ liệu đã sửa/xác minh. Tách bản gốc trong hồ sơ mới để giữ lịch sử.')
        original = PdfReader(doc['path'])
        if original.is_encrypted or len(original.pages) > reader.MAX_PAGES:
            raise HTTPException(422, 'PDF khóa hoặc vượt giới hạn trang.')
        pdf_bundle.validate_ranges(body.groups, len(original.pages))
        pages = db.rows('SELECT * FROM pages WHERE document_id=? ORDER BY number', (id,))
        if len(pages) != len(original.pages):
            raise HTTPException(409, 'Đọc/OCR hết PDF trước khi tách để có căn cứ theo trang.')
        try:
            with db.connect() as c:
                source_id = db.uid()
                c.execute('INSERT INTO bundle_sources VALUES (?,?,?,?,?)', (source_id, doc['case_id'], doc['name'], doc['hash'], doc['path']))
                # Transaction replaces the working bundle; the physical original is retained.
                c.execute('DELETE FROM documents WHERE id=?', (id,))
                for index, group in enumerate(body.groups):
                    writer = PdfWriter()
                    for n in range(group.start - 1, group.end):
                        writer.add_page(original.pages[n])
                    buffer = io.BytesIO()
                    writer.write(buffer)
                    data = buffer.getvalue()
                    digest = hashlib.sha256(data).hexdigest()
                    if c.execute('SELECT id FROM documents WHERE case_id=? AND hash=?', (doc['case_id'], digest)).fetchone():
                        raise ValueError('Một phần PDF trùng nội dung đã có. Chưa tách; file gốc giữ nguyên.')
                    child_id = db.uid()
                    path = db.DATA / 'files' / (child_id + '.pdf')
                    path.write_bytes(data)
                    created_paths.append(path)
                    name = f"{Path(doc['name']).stem[:85]} - {index+1:02d} {group.type.replace('/', '-')} (p{group.start}-{group.end}).pdf"
                    c.execute('INSERT INTO documents(id,case_id,name,ext,size,hash,path,status,type,reading_method) VALUES (?,?,?,?,?,?,?,?,?,?)', (child_id, doc['case_id'], name, '.pdf', len(data), digest, str(path), 'read', group.type, doc['reading_method']))
                    c.execute('INSERT INTO document_sources VALUES (?,?,?,?)', (child_id, source_id, group.start, group.end))
                    child_pages = pages[group.start-1:group.end]
                    for n, p in enumerate(child_pages, 1):
                        c.execute('INSERT INTO pages VALUES (?,?,?,?,?,?,?)', (db.uid(), child_id, n, p['text'], p['method'], p['quality'], p['table_json']))
                    # Extraction stays pending and classification still requires staff confirmation.
                    values = [{**p, 'number': n, 'tables': json.loads(p['table_json'])} for n, p in enumerate(child_pages, 1)]
                    for record in extraction.extract(values, group.type):
                        add_record(c, doc['case_id'], child_id, record)
                db.touch(c, doc['case_id'])
                db.audit(c, doc['case_id'], 'Tách bộ PDF; giữ bản gốc', source_id, before={'name': doc['name'], 'hash': doc['hash']}, after={'groups': [g.model_dump() for g in body.groups]})
                revalidate(c, doc['case_id'])
        except Exception:
            for path in created_paths:
                path.unlink(missing_ok=True)
            raise
    return bundle(doc['case_id'])

@app.get('/api/documents/{id}/pages/{number}')
def get_page(id: str, number: int):
    require_doc(id)
    page = db.one('SELECT number,text,method,quality,table_json FROM pages WHERE document_id=? AND number=?', (id, number))
    if not page:
        raise HTTPException(404, 'Trang chưa được đọc')
    page['tables'] = json.loads(page.pop('table_json'))
    return page

@app.get('/api/documents/{id}/preview/{number}')
def preview(id: str, number: int):
    doc = require_doc(id)
    if number < 1 or number > reader.MAX_PAGES:
        raise HTTPException(404, 'Trang không hợp lệ')
    if doc['ext'] == '.pdf':
        try:
            png_bytes = reader.render_pdf_page_png(doc['path'], number, scale=1.5)
            return Response(png_bytes, media_type='image/png')
        except ValueError as exc:
            raise HTTPException(404, str(exc))
    if doc['ext'] in ['.png', '.jpg', '.jpeg', '.tif', '.tiff']:
        with Image.open(doc['path']) as im:
            if number > getattr(im, 'n_frames', 1):
                raise HTTPException(404, 'Không có trang này')
            im.seek(number - 1)
            im = ImageOps.exif_transpose(im.copy())
            im.thumbnail((1800, 2200))
            buf = io.BytesIO()
            im.convert('RGB').save(buf, format='PNG')
            return Response(buf.getvalue(), media_type='image/png')
    raise HTTPException(404, 'Dùng xem văn bản/bảng cho định dạng này')

@app.get('/api/records/{id}/source')
def record_source(id: str):
    record=require_record(id)
    if not record['document_id'] or not record['page']:
        raise HTTPException(422,'Bản ghi chưa có chứng từ/trang nguồn. Bổ sung căn cứ trước khi xác minh.')
    doc=require_working_doc(record['document_id'])
    page=db.one('SELECT number,text,method,quality FROM pages WHERE document_id=? AND number=?',(doc['id'],record['page']))
    if not page:
        raise HTTPException(404,'Trang nguồn chưa được đọc. Chạy Đọc & trích xuất trước.')
    if any(j['status']=='running' for j in JOBS.values()):
        raise HTTPException(409,'Đang đọc chứng từ. Chờ hoàn tất rồi mở vùng nguồn để tránh quá tải OCR.')
    return source_locator.inspect(doc,record,page)

def document_diagnostic(doc,index):
    entries=json.loads(doc.get('reading_diagnostics') or '[]')
    if index<0 or index>=len(entries):raise HTTPException(404,'Mục đọc lại không còn tồn tại; tải lại hồ sơ.')
    return entries,entries[index]

@app.get('/api/documents/{id}/review-source')
def diagnostic_source(id: str, diagnostic: Optional[int]=None, number: int=1,field_id:Optional[str]=None):
    doc=require_working_doc(id)
    raw='';evidence=doc.get('classification_evidence','')
    if not evidence and diagnostic is None and not field_id:
        source_pages=db.rows('SELECT number,text FROM pages WHERE document_id=? ORDER BY number',(id,))
        _,evidence=extraction.classify(source_pages)
        if evidence:
            number=next((p['number'] for p in source_pages if evidence in p['text']),number)
    if diagnostic is not None:
        _,entry=document_diagnostic(doc,diagnostic)
        number=entry.get('page') or number
        candidate=entry.get('candidate_raw',{})
        raw=candidate.get('raw_value','') if entry.get('candidate_kind')=='field' else candidate
        evidence=entry.get('evidence','')
        field_id=entry.get('field_id')
    elif field_id:
        if field_id not in REVIEW_FIELDS:raise HTTPException(422,'Trường dữ liệu không thuộc phạm vi.')
        import re
        aliases=FIELDS[field_id]['aliases']
        matches=[(p['number'],line) for p in db.rows('SELECT number,text FROM pages WHERE document_id=? ORDER BY number',(id,))
                 for line in p['text'].splitlines() if any(re.search(r'(?i)(?<!\w)'+re.escape(a)+r'(?!\w)',line) for a in aliases)]
        if matches:number,evidence=matches[0]
    page=db.one('SELECT number,text,method,quality FROM pages WHERE document_id=? AND number=?',(id,number))
    if not page:raise HTTPException(404,'Trang nguồn chưa đọc được; cần đọc lại tài liệu.')
    busy(doc['case_id'])
    result=source_locator.inspect(doc,{'key':field_id,'raw':raw,'evidence':evidence,'method':'OpenAI diagnostic'},page)
    result['message']=('Vùng đánh dấu là chữ/đoạn trích dẫn có thật để nhân viên kiểm tra; giá trị bị loại chưa được xác nhận. '+result['message']).strip()
    return result

@app.post('/api/documents/{id}/diagnostics/{index}/review')
def review_diagnostic(id: str,index: int,body: DiagnosticReviewInput):
    with LOCK:
        doc=require_working_doc(id);busy(doc['case_id'])
        entries,entry=document_diagnostic(doc,index)
        if entry.get('review') or body.token!=source_locator.diagnostic_token(entry):
            raise HTTPException(409,'Kết quả đã thay đổi/đã xử lý. Tải lại trước khi xác nhận.')
        note=body.note.strip()
        if not body.confirmed and len(note)<5:raise HTTPException(422,'Đánh dấu đã kiểm tra bản gốc trước khi lưu.')
        note=note or ('Nhân viên đã kiểm tra bản gốc và xác nhận giá trị.' if body.action=='accept' else 'Nhân viên đã kiểm tra bản gốc và loại đề xuất.')
        page=entry.get('page')
        if not page or not db.one('SELECT id FROM pages WHERE document_id=? AND number=?',(id,page)):
            raise HTTPException(422,'Thiếu trang nguồn; đọc lại tệp trước khi xác nhận.')
        rid=None
        with db.connect() as c:
            if body.action=='accept':
                kind=entry.get('candidate_kind');key=entry.get('field_id') if kind=='field' else 'goods_item'
                if kind not in ['field','item'] or (kind=='field' and key not in REVIEW_FIELDS):
                    raise HTTPException(422,'Lần đọc cũ thiếu cấu trúc dữ liệu. Đọc lại bằng OpenAI trước.')
                value=checked_value(kind,key,body.value)
                if value.get('warning'):raise HTTPException(422,'Giá trị chưa rõ định dạng; sửa số/đơn vị trước khi xác nhận.')
                if kind=='item':
                    from .normalize import decimal_text
                    if decimal_text(value.get('quantity','')) is None or not value.get('unit'):
                        raise HTTPException(422,'Dòng hàng cần số lượng và đơn vị rõ ràng.')
                    if value.get('scope','invoice')=='invoice' and any(decimal_text(value.get(k,'')) is None for k in ['unit_price','line_amount']):
                        raise HTTPException(422,'Dòng hóa đơn cần đơn giá và thành tiền rõ ràng.')
                rid=db.uid()
                c.execute('INSERT INTO records VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                    (rid,doc['case_id'],id,kind,key,db.encode(entry.get('candidate_raw',{})),db.encode(value),page,
                     (entry.get('evidence','')+'\n[Nhân viên kiểm tra] '+note).strip(),'Manual review','verified',note,db.now(),1))
            entry['review']={'action':body.action,'note':note,'at':db.now(),'record_id':rid}
            c.execute('UPDATE documents SET reading_diagnostics=? WHERE id=?',(db.encode(entries),id))
            db.touch(c,doc['case_id'])
            db.audit(c,doc['case_id'],'Kiểm tra dữ liệu AI bị loại',id+':'+str(index),None,entry['review'],note)
            revalidate(c,doc['case_id'])
    return bundle(doc['case_id'])

@app.post('/api/checks/{id}/review')
def review_check(id: str,body: CheckReviewInput):
    with LOCK:
        check=db.one('SELECT * FROM checks WHERE id=?',(id,))
        if not check:raise HTTPException(404,'Kết quả đã thay đổi. Mở lại mục cần kiểm tra.')
        busy(check['case_id'])
        if check['key'].startswith(('reading:','classification:')) or check['result']=='Missing':
            raise HTTPException(422,'Mục này cần bổ sung/kiểm tra từng dữ liệu nguồn.')
        ids={s['id'] for s in json.loads(check['detail_json']).get('sources',[])}
        rows=[r for r in working_records(check['case_id']) if r['id'] in ids and r['status']!='rejected']
        if not rows or body.versions!={r['id']:r['version'] for r in rows}:
            raise HTTPException(409,'Phiên bản dữ liệu đã thay đổi hoặc không có bản ghi để xác nhận.')
        from .acceptance import readable
        if any(not readable(r) for r in rows):raise HTTPException(422,'Nhập lại ô còn trống hoặc chưa đọc rõ trước khi xác nhận.')
        if not body.confirmed and len(body.note.strip())<5:raise HTTPException(422,'Đánh dấu đã kiểm tra bản gốc trước khi lưu.')
        note=body.note.strip() or 'Nhân viên đã đối chiếu các nguồn và xác nhận đã kiểm tra.'
        with db.connect() as c:
            for r in rows:
                c.execute("UPDATE records SET status='verified',note=?,updated_at=?,version=version+1 WHERE id=?",(note,db.now(),r['id']))
            c.execute('UPDATE checks SET resolution=? WHERE id=?',(note,id))
            db.touch(c,check['case_id'])
            db.audit(c,check['case_id'],'Nhân viên xác nhận đã kiểm tra nguồn',id,None,[r['id'] for r in rows],note)
    return bundle(check['case_id'])

def worker(job_id):
    job = JOBS[job_id]
    id = job['case_id']
    mode = job['mode']
    final_status = 'error'
    try:
        docs = db.rows("SELECT * FROM documents WHERE case_id=? AND collection='document' ORDER BY rowid", (id,))
        if job.get('document_id'):
            docs = [d for d in docs if d['id']==job['document_id']]
        for idx, d in enumerate(docs):
            if job['cancel']:
                raise InterruptedError()
            job.update(document=d['name'], index=idx + 1, total=len(docs), stage='Đọc / OCR', pages=0)
            started = time.monotonic()
            try:
                uncertain_image = False
                reader_signature = pipeline.fingerprint('reader')
                reused_pages = d['status'] in ['read', 'needs_split'] and d.get('reader_signature')==reader_signature
                if mode == 'openai' and d['ext'] in {'.png', '.jpg', '.jpeg', '.tif', '.tiff'} and not d['type_verified']:
                    job['stage'] = 'Kiểm tra ảnh có chứa chứng từ'
                    route = ai.route_image(id, d)
                    if job['cancel']:
                        raise InterruptedError()
                    if route.category == 'reference' and route.certainty == 'clear':
                        with db.connect() as c:
                            move_collection(c, d, 'reference', route.reason)
                        job['warnings'].append({'document': d['name'], 'message': 'Đã đưa vào Giấy tờ khác: ' + route.reason + '. Có thể kiểm bản gốc và đưa lại vào bộ chính.'})
                        continue
                    uncertain_image = route.category == 'uncertain' or route.certainty != 'clear'
                if reused_pages:
                    pages = db.rows('SELECT number,text,method,quality,table_json FROM pages WHERE document_id=? ORDER BY number', (d['id'],))
                    for p in pages:
                        p['tables'] = json.loads(p.pop('table_json'))
                    if mode=='openai' and any(p['method']=='OCR' and 'Đã đọc bổ sung ô/dòng theo bố cục.' not in p['quality'] for p in pages):
                        pages=reader.read_document(d['path'],lambda n:job.update(pages=n),lambda:job['cancel'])
                        with db.connect() as c:
                            c.execute('DELETE FROM pages WHERE document_id=?',(d['id'],))
                            for p in pages:c.execute('INSERT INTO pages VALUES (?,?,?,?,?,?,?)',(db.uid(),d['id'],p['number'],p['text'],p['method'],p['quality'],db.encode(p['tables'])))
                else:
                    db.execute("UPDATE documents SET status='reading',error=NULL WHERE id=?", (d['id'],))
                    pages = reader.read_document(d['path'], lambda n: job.update(pages=n), lambda: job['cancel'])
                    with db.connect() as c:
                        c.execute('DELETE FROM pages WHERE document_id=?', (d['id'],))
                        for p in pages:
                            c.execute('INSERT INTO pages VALUES (?,?,?,?,?,?,?)', (db.uid(), d['id'], p['number'], p['text'], p['method'], p['quality'], db.encode(p['tables'])))
                        c.execute(
                            "UPDATE documents SET status='read',reading_method=?,elapsed=?,reader_signature=? WHERE id=?",
                            (', '.join(sorted({p['method'] for p in pages})), time.monotonic() - started, reader_signature, d['id']),
                        )
                job['stage'] = 'Phân loại / trích xuất ' + ('Local' if mode == 'local' else 'OpenAI')
                extraction_warning = None
                reading_diagnostics = []
                source_grounded = False
                extraction_trace = {}
                if d['ext'] == '.pdf' and pdf_bundle.is_mixed(pages):
                    message = 'Đã đọc PDF gồm nhiều loại chứng từ. Chọn “Tách bộ PDF” để phân bổ trang trước khi trích xuất; bản gốc được giữ lại.'
                    with db.connect() as c:
                        c.execute("UPDATE documents SET status='needs_split',error=NULL,type_verified=0 WHERE id=?", (d['id'],))
                        db.audit(c, id, 'Đọc bộ PDF; chờ tách theo trang', d['id'], after={'pages':len(pages),'groups':pdf_bundle.suggest(pages)})
                    job['warnings'].append({'document':d['name'],'message':message})
                    continue
                if uncertain_image:
                    kind, evidence, records = 'Other', route.reason, []
                    extraction_warning = 'Chưa rõ ảnh có chứa chứng từ. Kiểm tra bản gốc, xác nhận đúng loại chứng từ rồi đọc lại; chưa trích xuất dữ liệu.'
                    job['warnings'].append({'document': d['name'], 'message': extraction_warning})
                elif mode == 'openai':
                    result, records, info = ai.analyze(id, d, pages)
                    source_grounded = info.get('source_grounded',False)
                    extraction_trace = {**info.get('trace',{}), 'cache_source':info.get('source','unknown')}
                    kind = result.document_type
                    evidence = result.classification_evidence
                    if info['rejected']:
                        extraction_warning = ai.extraction_notice(info['rejected'])
                        reading_diagnostics = ai.extraction_diagnostics(result, pages, info['rejected'])
                        if extraction_warning:
                            job['warnings'].append({'document': d['name'], 'message': extraction_warning})
                else:
                    kind, evidence = extraction.classify(pages)
                    if d['type_verified'] and not (d['type']=='Contract' and kind=='Purchase Order'):
                        kind = d['type']
                    records = extraction.extract(pages, kind)
                if job['cancel']:
                    raise InterruptedError()
                case = require_case(id)
                commercial = commercial_profile(json.loads(case['requirements']),case['mode'])
                main_types = COMPARISON_TYPES
                # A staff correction is explicit. Keep it in the main set, but flag conflicting AI classification.
                if commercial and d['type_verified'] and d['type'] in main_types and kind != d['type'] and not (d['type']=='Contract' and kind=='Purchase Order'):
                    extraction_warning = f'Máy gợi ý {kind}, khác loại {d["type"]} đã được nhân viên xác nhận. Kiểm bản gốc và loại trước khi dùng dữ liệu.'
                    kind, records = d['type'], []
                    job['warnings'].append({'document':d['name'],'message':extraction_warning})
                if commercial and kind not in main_types and not uncertain_image:
                    reason = ('Không đủ căn cứ nhận diện 6 loại chính; cần xem lại bản gốc.' if kind == 'Other'
                              else f'Nhận diện {kind}; nằm ngoài 6 loại chứng từ đang kiểm thử.')
                    with db.connect() as c:
                        c.execute("UPDATE documents SET type=?,classification_evidence=?,status='read',error=NULL WHERE id=?",(kind,evidence,d['id']))
                        move_collection(c,d,'reference',reason)
                        db.audit(c,id,'Phân loại vào Giấy tờ khác',d['id'],after={'type':kind,'evidence':evidence,'method':mode},reason=reason)
                    job['warnings'].append({'document':d['name'],'message':'Giấy tờ khác: '+reason})
                    continue
                if job['cancel']:
                    raise InterruptedError()
                with db.connect() as c:
                    c.execute(
                        'UPDATE documents SET type=?,classification_evidence=?,type_verified=? WHERE id=?',
                        (kind, evidence, d['type_verified'] if kind == d['type'] else 0, d['id']),
                    )
                    replace_extracted_records(c, id, d['id'], [r for r in records if not (commercial and kind in ['Arrival Notice','B/L','Sea Waybill','AWB','Customs Declaration','Customs Release','Warehouse Receipt'] and r['kind']!='field')])
                    c.execute("UPDATE documents SET status='read',error=?,reading_diagnostics=?,extraction_required=0,classification_method=? WHERE id=?", (extraction_warning, db.encode(reading_diagnostics), 'OpenAI' if mode=='openai' else 'Local', d['id']))
                    db.audit(c, id, 'Đọc và trích xuất ' + mode, d['id'], after={'type': kind, 'pages': len(pages), 'records': len(records),'source_format_grounded':source_grounded,
                        'trace':{**extraction_trace,'pages_sha256':pipeline.digest(pages),'reader_sha256':reader_signature,'pages_source':'saved' if reused_pages else 'fresh'}})
            except InterruptedError:
                raise
            except Exception as exc:
                message = str(exc)[:600] if isinstance(exc, ValueError) else 'Không đọc/xử lý được tệp. Kiểm tra định dạng và thử lại.'
                db.execute("UPDATE documents SET status='error',error=? WHERE id=?", (message, d['id']))
                job['errors'].append({'document': d['name'], 'error': message})
            finally:
                reader.release_reader_memory()
        if mode=='openai' and not job['cancel'] and commercial_profile(json.loads(require_case(id)['requirements']),require_case(id)['mode']):
            job['stage']='OpenAI đối chiếu toàn bộ hồ sơ'
            case_docs=db.rows("SELECT * FROM documents WHERE case_id=? AND collection='document' ORDER BY rowid",(id,))
            clause_pages=db.rows("SELECT p.document_id,p.number,p.text FROM pages p JOIN documents d ON d.id=p.document_id WHERE d.case_id=? AND d.collection='document' AND d.type IN ('Contract','Purchase Order') ORDER BY d.id,p.number",(id,))
            review_error=consistency.run_ai_review(id,case_docs,working_records(id),clause_pages,True)
            if review_error:job['warnings'].append({'document':'Đối chiếu toàn hồ sơ','message':review_error})
        final_status = 'completed'
        if job['errors']:
            job['stage'] = f"Hoàn tất có cảnh báo ({len(job['errors'])}/{len(docs)} tệp lỗi — xem chi tiết từng tệp)"
        elif job['warnings']:
            job['stage'] = f"Đã đọc; {len(job['warnings'])} cảnh báo cần kiểm tra"
        else:
            job['stage'] = 'Đã đối chiếu; mở Sai lệch & xử lý'
    except InterruptedError:
        final_status = 'cancelled'
        job['stage'] = 'Đã hủy'
        db.execute("UPDATE documents SET status='error',error='Đã hủy xử lý. Chọn thử lại.' WHERE case_id=? AND status='reading'", (id,))
    except Exception:
        final_status = 'error'
        job['stage'] = 'Lỗi xử lý; thử lại'
    finally:
        with db.connect() as c:
            db.touch(c, id)
            revalidate(c, id)
        job['status'] = final_status
        job['ended_at'] = db.now()

@app.post('/api/cases/{id}/process')
def process(id: str, body: ProcessInput):
    return start_processing(id,body)

@app.post('/api/documents/{id}/process')
def process_document(id: str, body: ProcessInput):
    doc = require_working_doc(id)
    return start_processing(doc['case_id'],body,id)

def start_processing(id: str, body: ProcessInput, document_id=None):
    with LOCK:
        require_case(id)
        busy(id)
        if document_id:
            selected = require_working_doc(document_id)
            if selected['case_id'] != id:
                raise HTTPException(422, 'Tệp không thuộc hồ sơ đang xử lý.')
        if os.getenv('UI_SMART_CLOUD') == '1' and any(j['status'] == 'running' for j in JOBS.values()):
            raise HTTPException(409, 'Máy chủ đang xử lý hồ sơ khác. Chờ hoàn tất rồi thử lại để tránh quá tải OCR.')
        if body.mode == 'openai' and not ai.key():
            raise HTTPException(409, 'Chưa cấu hình OpenAI. Tiếp tục Local hoặc mở Cài đặt.')
        if len(JOBS) >= 200:
            for old_id in list(JOBS):
                if JOBS[old_id]['status'] != 'running':
                    del JOBS[old_id]
                if len(JOBS) < 150:
                    break
        if body.mode == 'openai' and not body.allow_external:
            raise HTTPException(422, 'Cần chủ động đồng ý gửi nội dung chứng từ ra OpenAI cho yêu cầu này.')
        if not db.one("SELECT id FROM documents WHERE case_id=? AND collection='document'", (id,)):
            raise HTTPException(422, 'Tải ít nhất một chứng từ')
        job_id = db.uid()
        JOBS[job_id] = {
            'id': job_id,
            'case_id': id,
            'document_id': document_id,
            'mode': body.mode,
            'status': 'running',
            'stage': 'Tiếp nhận',
            'document': '',
            'index': 0,
            'total': 0,
            'pages': 0,
            'cancel': False,
            'errors': [],
            'warnings': [],
            'started_at': db.now(),
        }
        db.execute("UPDATE cases SET status='Reading' WHERE id=?", (id,))
        threading.Thread(target=worker, args=(job_id,), daemon=True).start()
    return JOBS[job_id]

@app.get('/api/jobs/{id}')
def job_status(id: str):
    if id not in JOBS:
        raise HTTPException(404, 'Tác vụ không còn trong phiên hiện tại')
    return JOBS[id]

@app.get('/api/samples/{name}')
def sample_document(name: str):
    # Only synthetic fixtures explicitly shipped for the two UI demo sets.
    if name not in {'invoice.txt', 'packing.txt', 'packing-conflict.txt', 'arrival.txt', 'booking.txt', 'contract.txt', 'po.txt', 'co.txt'}:
        raise HTTPException(404, 'Không có chứng từ mẫu này')
    return FileResponse(db.ROOT / 'samples' / name, filename=name, media_type='application/pdf' if name.endswith('.pdf') else 'text/plain')

@app.post('/api/jobs/{id}/cancel')
def cancel(id: str):
    job = job_status(id)
    job['cancel'] = True
    return {'ok': True, 'detail': 'Dừng sau trang/yêu cầu hiện tại; OCR timeout tối đa 60 giây.'}

def checked_value(kind, key, value):
    if kind == 'field':
        if key not in FIELDS or not isinstance(value, str):
            raise ValueError('Field ID/giá trị không hợp lệ')
        if not value.strip() or len(value) > 2000:
            raise ValueError('Giá trị cần từ 1 đến 2000 ký tự; dùng Từ chối để bỏ bản ghi sai')
        return normalize(value, FIELDS[key]['type'])
    if not isinstance(value, dict):
        raise ValueError('Dòng hàng/container cần object')
    allowed = (set(extraction.ITEM_ALIASES) | {'scope', 'warning'}) if kind == 'item' else set(extraction.CONT_ALIASES) | {'warning'}
    if set(value) - allowed:
        raise ValueError('Có cột chưa hỗ trợ')
    if any(not isinstance(v, (str, type(None))) or len(v or '') > 1000 for v in value.values()):
        raise ValueError('Giá trị dòng quá dài/không hợp lệ')
    if kind == 'item' and not value.get('description'):
        raise ValueError('Cần mô tả hàng')
    if kind == 'item':
        from .normalize import decimal_text
        import re
        # Recompute warnings after a staff edit; never retain a stale OCR warning
        # or allow a caller to clear an invalid number by removing the warning flag.
        value = {k:v for k,v in value.items() if k != 'warning'}
        if value.get('line_number') and not re.fullmatch(r'\d+', value['line_number']):
            value['warning'] = 'Số thứ tự dòng chưa rõ; đối chiếu bảng gốc và sửa trước khi xác minh.'
        if any(value.get(k) and decimal_text(value[k]) is None for k in ['quantity','unit_price','line_amount']):
            value['warning'] = 'Có số chưa rõ định dạng; cần sửa trước khi xác minh.'
    if kind == 'container':
        import re
        value = {**value, 'container_number': validation.container_number(value.get('container_number'))}
        if not re.fullmatch('[A-Z]{4}[0-9]{7}', value.get('container_number', '')):
            raise ValueError('Container cần 4 chữ + 7 số')
    return value

@app.post('/api/cases/{id}/records')
def manual(id: str, body: NewRecord):
    with LOCK:
        case=require_case(id)
        busy(id)
        if commercial_profile(json.loads(case['requirements']),case['mode']) and (body.kind == 'container' or (body.kind == 'field' and body.key not in REVIEW_FIELDS)):
            raise HTTPException(422, 'Chỉ thêm trường chứng từ hoặc dòng hàng trong phạm vi checklist hiện tại')
        value = checked_value(body.kind, body.key, body.value)
        if body.confirmed:
            from .acceptance import readable
            if not readable({'kind':body.kind,'key':body.key,'value':value}):
                raise HTTPException(422,'Nhập đầy đủ giá trị và đơn vị đọc được trước khi xác nhận.')
        if body.document_id:
            d = require_working_doc(body.document_id)
            if d['case_id'] != id:
                raise HTTPException(422, 'Chứng từ thuộc hồ sơ khác')
            if body.page and not db.one('SELECT id FROM pages WHERE document_id=? AND number=?', (body.document_id, body.page)):
                raise HTTPException(422, 'Trang nguồn không tồn tại')
        if commercial_profile(json.loads(case['requirements']),case['mode']) and (not body.document_id or not body.page):
            raise HTTPException(422, 'Chọn chứng từ và trang nguồn chứa giá trị cần bổ sung')
        with db.connect() as c:
            add_record(c, id, body.document_id, {'kind': body.kind, 'key': body.key, 'raw': body.value, 'value': value, 'page': body.page, 'evidence': body.note, 'method': 'Manual'})
            if body.confirmed:
                c.execute("UPDATE records SET status='verified' WHERE rowid=last_insert_rowid()")
            db.touch(c, id)
            db.audit(c, id, 'Thêm dữ liệu thủ công', after=body.model_dump())
            revalidate(c, id)
    return bundle(id)

@app.patch('/api/records/{id}')
def edit(id: str, body: EditInput):
    with LOCK:
        r = require_record(id)
        busy(r['case_id'])
        if r['version'] != body.version:
            raise HTTPException(409, 'Dữ liệu đã thay đổi; tải lại trước khi sửa')
        value = checked_value(r['kind'], r['key'], body.value)
        with db.connect() as c:
            c.execute('INSERT INTO revisions VALUES (?,?,?,?,?,?,?)', (db.uid(), id, db.encode(r['value']), db.encode(value), 'Nhân viên local', body.note, db.now()))
            c.execute(
                "UPDATE records SET value_json=?,status='pending',method='Manual edit',note=?,updated_at=?,version=version+1 WHERE id=?",
                (db.encode(value), body.note, db.now(), id),
            )
            db.touch(c, r['case_id'])
            db.audit(c, r['case_id'], 'Sửa dữ liệu', id, r['value'], value, body.note)
            revalidate(c, r['case_id'])
    return bundle(r['case_id'])

@app.post('/api/records/{id}/verify')
def verify(id: str, body: VerifyInput):
    with LOCK:
        r = require_record(id)
        busy(r['case_id'])
        if r['version'] != body.version:
            raise HTTPException(409, 'Dữ liệu đã thay đổi; tải lại')
        if body.status == 'verified':
            if r['value'].get('warning'):
                raise HTTPException(422, 'Cần sửa dữ liệu chưa rõ định dạng trước khi xác minh')
            if not r['evidence'] and not body.note:
                raise HTTPException(422, 'Cần bằng chứng/ghi chú xác minh')
            if r['key'] in ['hs_code', 'tax_rate', 'customs_value'] and len(body.note) < 5:
                raise HTTPException(422, 'Chỉ tiêu pháp lý cần ghi rõ căn cứ xác minh')
            if r['kind'] == 'item':
                from .normalize import decimal_text
                v = r['value']
                if decimal_text(v.get('quantity', '')) is None or not v.get('unit'):
                    raise HTTPException(422, 'Dòng hàng cần số lượng và đơn vị rõ ràng')
                if v.get('scope') != 'packing' and any(decimal_text(v.get(k, '')) is None for k in ['unit_price', 'line_amount']):
                    raise HTTPException(422, 'Dòng Invoice cần đơn giá và thành tiền rõ ràng')
        with db.connect() as c:
            c.execute('UPDATE records SET status=?,note=?,updated_at=?,version=version+1 WHERE id=?', (body.status, body.note or r['note'], db.now(), id))
            db.touch(c, r['case_id'])
            db.audit(c, r['case_id'], 'Xác minh từng bản ghi', id, r['status'], body.status, body.note)
            revalidate(c, r['case_id'])
    return bundle(r['case_id'])

@app.post('/api/records/{id}/undo')
def undo(id: str):
    with LOCK:
        r = require_record(id)
        busy(r['case_id'])
        rev = db.one('SELECT * FROM revisions WHERE record_id=? ORDER BY rowid DESC LIMIT 1', (id,))
        if not rev:
            raise HTTPException(409, 'Không có chỉnh sửa để hoàn tác')
        with db.connect() as c:
            c.execute("UPDATE records SET value_json=?,status='pending',updated_at=?,version=version+1 WHERE id=?", (rev['before_json'], db.now(), id))
            c.execute('DELETE FROM revisions WHERE id=?', (rev['id'],))
            db.touch(c, r['case_id'])
            db.audit(c, r['case_id'], 'Hoàn tác', id, r['value'], json.loads(rev['before_json']))
            revalidate(c, r['case_id'])
    return bundle(r['case_id'])

@app.post('/api/checks/{id}/resolve')
def resolve(id: str, body: ResolveInput):
    body.note = body.note.strip()
    if len(body.note) < 5:
        raise HTTPException(422, 'Ghi lý do xử lý rõ ràng, tối thiểu 5 ký tự không tính khoảng trắng.')
    with LOCK:
        check = db.one('SELECT * FROM checks WHERE id=?', (id,))
        if not check:
            raise HTTPException(404, 'Cảnh báo không tồn tại')
        busy(check['case_id'])
        if check['result'] == 'Missing':
            raise HTTPException(422, 'Bổ sung dữ liệu hoặc điều chỉnh checklist có lý do để xử lý mục thiếu')
        if check['key'].startswith(('reading:', 'classification:')):
            raise HTTPException(400, 'Cần đọc tài liệu và bấm Xác nhận loại chứng từ; ghi chú không thể bỏ qua bước này.')
        with db.connect() as c:
            c.execute('UPDATE checks SET resolution=? WHERE id=?', (body.note, id))
            db.touch(c, check['case_id'])
            db.audit(c, check['case_id'], 'Ghi nhận xử lý đối chiếu', id, check['resolution'], body.note, body.note)
    return bundle(check['case_id'])

@app.patch('/api/cases/{id}/requirements/{kind:path}')
def requirement(id: str, kind: str, body: RequirementInput):
    body.note = body.note.strip()
    if len(body.note) < 3:
        raise HTTPException(422, 'Checklist cần lý do rõ ràng, tối thiểu 3 ký tự không tính khoảng trắng.')
    with LOCK:
        case = require_case(id)
        busy(id)
        rules = json.loads(case['requirements'])
        if kind not in rules:
            raise HTTPException(422, 'Checklist không hỗ trợ loại này')
        before = rules[kind]
        rules[kind] = body.model_dump()
        with db.connect() as c:
            c.execute('UPDATE cases SET requirements=? WHERE id=?', (db.encode(rules), id))
            db.touch(c, id)
            db.audit(c, id, 'Đổi điều kiện checklist', kind, before, rules[kind], body.note)
            revalidate(c, id)
    return bundle(id)

@app.post('/api/cases/{id}/review')
def review(id: str):
    with LOCK:
        busy(id)
        b = bundle(id)
        if b['blockers']:
            raise HTTPException(409, {'message': 'Chưa đủ điều kiện Reviewed', 'blockers': b['blockers']})
        with db.connect() as c:
            c.execute("UPDATE cases SET status='Reviewed',reviewed_at=?,updated_at=? WHERE id=?", (db.now(), db.now(), id))
            db.audit(c, id, 'Hoàn thành kiểm tra hồ sơ')
    return bundle(id)

@app.get('/api/cases/{id}/audit')
def history(id: str):
    return db.rows('SELECT * FROM audit WHERE case_id=? ORDER BY rowid DESC', (id,))

@app.get('/api/audit')
def all_history():
    return db.rows('SELECT * FROM audit ORDER BY rowid DESC LIMIT 200')

@app.post('/api/cases/{id}/export')
def export(id: str, body: ExportInput):
    with LOCK:
        busy(id)
        b = bundle(id)
        if body.selected and (len(body.selected) > 100 or set(body.selected) - set(FIELDS)):
            raise HTTPException(422, 'Field xuất không hợp lệ')
        if body.mode == 'Reviewed' and (b['blockers'] or not b['case']['reviewed_at']):
            raise HTTPException(409, 'Cần kiểm tra từng trường và hoàn thành duyệt hồ sơ trước khi xuất Reviewed')
        if not b['records']:
            raise HTTPException(422, 'Chưa có dữ liệu để xuất; bạn có thể nhập thủ công')
        n = (db.one('SELECT max(version) n FROM exports WHERE case_id=?', (id,))['n'] or 0) + 1
        export_id = db.uid()
        path = db.DATA / 'exports' / (export_id + '.xlsx')
        stats = excel.build(path, b['case'], b['documents'], b['records'], b['checks'], history(id), body.mode, body.workbook_type, body.selected, n)
        with db.connect() as c:
            c.execute('INSERT INTO exports VALUES (?,?,?,?,?,?,?,?)', (export_id, id, db.now(), n, body.mode, body.workbook_type, str(path), b['case']['revision']))
            db.audit(c, id, 'Xuất Excel', export_id, after={'mode': body.mode, 'version': n, 'stats': stats})
            if body.mode == 'Reviewed':
                c.execute("UPDATE cases SET status='Exported',updated_at=? WHERE id=?", (db.now(), id))
    return {'id': export_id, 'download': '/api/exports/' + export_id, 'stats': stats}

@app.get('/api/exports/{id}')
def download(id: str):
    r = db.one('SELECT * FROM exports WHERE id=?', (id,))
    if not r:
        raise HTTPException(404, 'Không có phiên bản xuất này')
    return FileResponse(r['path'], filename=f"UI_SMART_{r['case_id']}_{r['mode']}_v{r['version']}.xlsx")

@app.get('/api/ai/status')
def ai_status():
    return ai.status()

@app.post('/api/ai/configure')
def configure(body: AIInput):
    ai.configure(body.api_key, body.model)
    return ai.status()

@app.post('/api/ai/test')
def test_ai():
    try:
        return ai.test_connection()
    except ValueError:
        raise
    except Exception as exc:
        raise HTTPException(422, ai.safe_error(exc)) from None

@app.delete('/api/ai/configure')
def remove_key():
    ai.remove()
    return ai.status()

@app.post('/api/backup')
def backup():
    with LOCK:
        if any(j['status'] == 'running' for j in JOBS.values()):
            raise HTTPException(409, 'Chờ tác vụ hoàn tất trước khi backup')
        id = db.uid()
        folder = db.DATA / 'backups'
        snapshot = folder / (id + '.sqlite')
        archive = folder / (id + '.zip')
        with db.connect() as source:
            target = sqlite3.connect(snapshot)
            source.backup(target)
            target.close()
        with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as z:
            z.write(snapshot, 'smart.sqlite')
            for sub in ['files', 'exports', 'templates']:
                for p in (db.DATA / sub).glob('*'):
                    if p.is_file():
                        z.write(p, sub + '/' + p.name)
        snapshot.unlink()
        return {'id': id, 'download': '/api/backups/' + id}

@app.get('/api/backups/{id}')
def get_backup(id: str):
    import re
    if not re.fullmatch('[a-f0-9]{32}', id):
        raise HTTPException(404)
    p = db.DATA / 'backups' / (id + '.zip')
    if not p.exists():
        raise HTTPException(404)
    return FileResponse(p, filename='UI_SMART_backup_' + id + '.zip')

@app.post('/api/backups/restore-empty')
async def restore_empty(file: UploadFile = File(...)):
    """Recover a restarted ephemeral server; never replace existing cases."""
    data=await file.read(100*1024**2+1)
    await file.close()
    if len(data)>100*1024**2 or Path(file.filename or '').suffix.lower()!='.zip':
        raise HTTPException(422,'Bản sao lưu cần là ZIP dưới 100 MB.')
    with LOCK:
        if any(j['status']=='running' for j in JOBS.values()) or db.one('SELECT count(*) n FROM cases')['n']:
            raise HTTPException(409,'Chỉ khôi phục khi máy chủ chưa có hồ sơ; dữ liệu hiện có được giữ nguyên.')
        archive=db.ROOT/'tmp'/('restore-upload-'+db.uid()+'.zip')
        archive.parent.mkdir(exist_ok=True)
        archive.write_bytes(data)
        try:
            from scripts.restore import restore
            restore(archive,db.DATA,check_server=False)
            for folder in ['files','exports','backups','templates']:(db.DATA/folder).mkdir(exist_ok=True)
            db.migrate()
            upgrade_transport_scope()
            refresh_staff_checks()
            return {'restored':db.one('SELECT count(*) n FROM cases')['n']}
        except (ValueError,zipfile.BadZipFile,sqlite3.Error) as exc:
            raise HTTPException(422,'Bản sao lưu không hợp lệ; không thể khôi phục.') from None
        finally:archive.unlink(missing_ok=True)

# Experimental generic template mapper, explicitly internal and not a customs software import certification.
@app.post('/api/templates')
async def template_upload(file: UploadFile = File(...)):
    data = await file.read(reader.MAX_BYTES + 1)
    await file.close()
    if len(data) > reader.MAX_BYTES or Path(file.filename or '').suffix.lower() != '.xlsx':
        raise HTTPException(422, 'Template cần là XLSX dưới 25 MB')
    id = db.uid()
    p = db.DATA / 'templates' / (id + '.xlsx')
    p.write_bytes(data)
    try:
        validate_zip(p)
        wb = load_workbook(p)
        sheets = {
            s.title: [
                {'cell': cell.coordinate, 'header': str(cell.value)}
                for row in s.iter_rows(min_row=1, max_row=min(10, s.max_row), max_col=min(50, s.max_column))
                for cell in row
                if cell.value is not None
            ]
            for s in wb
        }
        wb.close()
        db.execute('INSERT INTO templates VALUES (?,?,?,?,?)', (id, Path(file.filename).name, str(p), db.encode(sheets), '[]'))
    except Exception:
        p.unlink(missing_ok=True)
        raise ValueError('Không đọc được template XLSX')
    return {'id': id, 'name': Path(file.filename).name, 'sheets': sheets, 'scope': 'Mapping nội bộ; chưa xác minh tương thích phần mềm khai báo'}

@app.get('/api/templates')
def list_templates():
    rows = db.rows('SELECT id,name,sheets_json,mapping_json FROM templates')
    for r in rows:
        r['sheets'] = json.loads(r.pop('sheets_json'))
        r['mapping'] = json.loads(r.pop('mapping_json'))
    return rows

@app.post('/api/templates/{id}/export')
def template_export(id: str, body: TemplateInput):
    with LOCK:
        busy(body.case_id)
        case = require_case(body.case_id)
        template = db.one('SELECT * FROM templates WHERE id=?', (id,))
        if not template:
            raise HTTPException(404, 'Không tìm thấy template')
        if not body.mapping:
            raise HTTPException(422, 'Chọn ít nhất một mapping cụ thể')
        export_id = db.uid()
        p = db.DATA / 'exports' / (export_id + '.xlsx')
        count = templates.build(template, body.mapping, body.case_id, p)
        n = (db.one('SELECT max(version) n FROM exports WHERE case_id=?', (body.case_id,))['n'] or 0) + 1
        with db.connect() as c:
            c.execute('UPDATE templates SET mapping_json=? WHERE id=?', (db.encode(body.mapping), id))
            c.execute('INSERT INTO exports VALUES (?,?,?,?,?,?,?,?)', (export_id, body.case_id, db.now(), n, 'Draft', 'Template-internal', str(p), case['revision']))
            db.audit(c, body.case_id, 'Xuất theo mapping nội bộ', export_id, after={'template': id, 'cells': count})
    return {'id': export_id, 'download': '/api/exports/' + export_id, 'scope': 'Bản mapping nội bộ; chưa kiểm thử nhập phần mềm bên ngoài'}

def upgrade_commercial_checklists():
    """Apply the requested scope once; preserve files, review evidence and shared waivers."""
    with db.connect() as c:
        if c.execute('SELECT 1 FROM schema_versions WHERE version=2').fetchone():
            return
        for row in c.execute("SELECT * FROM cases WHERE mode='FCL'").fetchall():
            old=json.loads(row['requirements'])
            rules=default_requirements('FCL')
            for key in ['Invoice','Packing List']:
                if key in old:
                    rules[key]=old[key]
            c.execute('UPDATE cases SET requirements=? WHERE id=?',(db.encode(rules),row['id']))
            db.touch(c,row['id'])
            db.audit(c,row['id'],'Áp dụng checklist chứng từ thương mại',before=old,after=rules,reason='Bỏ yêu cầu B/L, container và khai báo; giữ chứng từ gốc')
            revalidate(c,row['id'])
        c.execute('INSERT INTO schema_versions VALUES(2,?)',(db.now(),))

upgrade_commercial_checklists()

def upgrade_arrival_notice_scope():
    """Retain source files/waivers; re-read arrival notices previously excluded by scope."""
    with db.connect() as c:
        if c.execute('SELECT 1 FROM schema_versions WHERE version=3').fetchone():
            return
        for row in c.execute("SELECT * FROM cases WHERE mode='FCL'").fetchall():
            old=json.loads(row['requirements'])
            if not commercial_profile(old,row['mode']):
                continue
            rules={**old}
            rules.setdefault('Arrival Notice',default_requirements()['Arrival Notice'])
            c.execute('UPDATE cases SET requirements=? WHERE id=?',(db.encode(rules),row['id']))
            for doc in c.execute("SELECT * FROM documents WHERE case_id=? AND type='Arrival Notice' AND collection='reference'",(row['id'],)).fetchall():
                c.execute("UPDATE documents SET collection='document',collection_reason='',extraction_required=1,error=NULL WHERE id=?",(doc['id'],))
                db.audit(c,row['id'],'Đưa Arrival Notice vào bộ chính',doc['id'],reason='Phạm vi mới; cần đọc lại theo nguồn trước khi xác minh')
            db.touch(c,row['id'])
            db.audit(c,row['id'],'Cập nhật 6 nhóm chứng từ',before=old,after=rules,reason='Thêm Arrival Notice; số Packing List không bắt buộc; B/L không thay Booking')
            revalidate(c,row['id'])
        c.execute('INSERT INTO schema_versions VALUES(3,?)',(db.now(),))

upgrade_arrival_notice_scope()


def upgrade_transport_scope():
    """Include the newly requested transport documents without losing waivers."""
    with db.connect() as c:
        if c.execute('SELECT 1 FROM schema_versions WHERE version=5').fetchone():return
        for row in c.execute("SELECT * FROM cases WHERE mode='FCL'").fetchall():
            rules=json.loads(row['requirements'])
            if not commercial_profile(rules,row['mode']):continue
            rules.setdefault('B/L',default_requirements()['B/L'])
            c.execute('UPDATE cases SET requirements=? WHERE id=?',(db.encode(rules),row['id']))
            for doc in c.execute("SELECT * FROM documents WHERE case_id=? AND collection='reference' AND type IN ('B/L','Sea Waybill','AWB','Customs Declaration','Customs Release')",(row['id'],)).fetchall():
                c.execute("UPDATE documents SET collection='document',collection_reason='',extraction_required=1,error=NULL WHERE id=?",(doc['id'],))
                db.audit(c,row['id'],'Đưa vận đơn / tờ khai vào bộ đối chiếu',doc['id'],reason='Phạm vi mới; cần đọc nguồn trước khi trích xuất')
            revalidate(c,row['id'])
        c.execute('INSERT INTO schema_versions VALUES(5,?)',(db.now(),))


upgrade_transport_scope()


def refresh_staff_checks():
    """Apply corrected comparison rules to saved cases without re-reading files."""
    with db.connect() as c:
        if c.execute('SELECT 1 FROM schema_versions WHERE version=4').fetchone():
            return
        for row in c.execute('SELECT id FROM cases').fetchall():
            revalidate(c,row['id'])
        c.execute('INSERT INTO schema_versions VALUES(4,?)',(db.now(),))


refresh_staff_checks()

if (db.ROOT / 'dist').exists():
    app.mount('/', StaticFiles(directory=db.ROOT / 'dist', html=True), name='frontend')
