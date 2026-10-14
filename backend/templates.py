"""User-confirmed scalar cell mappings. No claim of external software compatibility."""
import json, re
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from . import db, excel

def validate_mapping(template, mapping, case_id):
    if len(mapping) > 100:
        raise ValueError('Tối đa 100 ô mapping')
    wb = load_workbook(template['path'])
    resolved = []
    targets = set()
    try:
        if any(c.data_type == 'f' for s in wb for row in s for c in row):
            raise ValueError('Bản mapping này không xử lý template có công thức; dùng mẫu dữ liệu thuần.')
        docs = {d['id']: d['name'] for d in db.rows('SELECT id, name FROM documents WHERE case_id=?', (case_id,))}
        for m in mapping:
            if set(m) != {'sheet', 'cell', 'record_id'}:
                raise ValueError('Mapping cần sheet/cell/record_id')
            if m['sheet'] not in wb.sheetnames or not re.fullmatch(r'[A-Z]{1,2}[1-9][0-9]{0,3}', m['cell']):
                raise ValueError('Sheet/ô mapping không hợp lệ')
            cell = wb[m['sheet']][m['cell']]
            if isinstance(cell, MergedCell):
                raise ValueError('Không ghi vào ô phụ của vùng merged')
            target = (m['sheet'], m['cell'])
            if target in targets:
                raise ValueError('Hai trường cùng ghi một ô')
            targets.add(target)
            record = db.one("SELECT * FROM records WHERE id=? AND case_id=? AND kind='field' AND status!='rejected'", (m['record_id'], case_id))
            if not record:
                raise ValueError('Record không thuộc hồ sơ hoặc đã bị từ chối')
            record['document_name'] = docs.get(record['document_id'], 'Nhập tay')
            value = json.loads(record['value_json'])
            resolved.append((m, excel.value_text(value), record))
        return wb, resolved
    except Exception:
        wb.close()
        raise

def build(template, mapping, case_id, path):
    wb, resolved = validate_mapping(template, mapping, case_id)
    try:
        for m, v, _ in resolved:
            cell = wb[m['sheet']][m['cell']]
            cell.value = excel.safe(v)
            cell.data_type = 's'
            cell.number_format = '@'
        # Explicit trace sheet makes unverified mapped data visible in a working copy.
        name = 'UI_SMART_NGUON'
        if name in wb.sheetnames:
            del wb[name]
        excel.sheet(
            wb,
            name,
            ['Ô đích', 'Record ID', 'Chỉ tiêu', 'Giá trị', 'Trạng thái', 'Nguồn', 'Trang', 'Bằng chứng'],
            [[m['sheet'] + '!' + m['cell'], r['id'], r['key'], v, r['status'], r.get('document_name') or r['document_id'] or 'Nhập tay', r['page'], r['evidence']] for m, v, r in resolved],
        )
        wb.save(path)
        check = load_workbook(path)
        try:
            for m, v, _ in resolved:
                assert check[m['sheet']][m['cell']].value == excel.safe(v)
        finally:
            check.close()
        return len(resolved)
    finally:
        wb.close()
